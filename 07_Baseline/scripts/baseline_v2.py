#!/usr/bin/env python3
"""
baseline_v2.py - Clinical Baseline Characteristics Table 1
Compares nonRMPP (label=0) vs RMPP (label=1) from an SPSS .sav file.
Reads SAV path and Excel output path from baseline_config.json.
"""

import json, os, warnings, sys
import numpy as np
import pandas as pd
import pyreadstat
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# ---------------------------------------------------------------------------
# 0. Load config
# ---------------------------------------------------------------------------
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(SCRIPT_DIR, "baseline_config.json"), "r", encoding="utf-8") as f:
    cfg = json.load(f)

SAV_PATH = cfg["sav"]
OUT_PATH = cfg["out"]
print(f"[INFO] Loading SAV: {SAV_PATH}")

# ---------------------------------------------------------------------------
# 1. Read SPSS file and split groups
# ---------------------------------------------------------------------------
df, meta = pyreadstat.read_sav(SAV_PATH)
df_non = df[df["label"] == 0.0].copy()
df_rmpp = df[df["label"] == 1.0].copy()
n_non = len(df_non)
n_rmpp = len(df_rmpp)
print(f"[INFO] nonRMPP (label=0): n={n_non}  |  RMPP (label=1): n={n_rmpp}")

# ---------------------------------------------------------------------------
# 2. Variable definitions
# ---------------------------------------------------------------------------
NORMAL_VARS = [
    "Age", "Weight", "PreadmissionFeverDuration", "LengthOfStay",
    "WBC", "HB", "PLT", "PT", "APTT",
]
NON_NORMAL_VARS = [
    "CRP", "NEUT", "LYMPH", "ALT", "AST", "LDH", "CK", "CKMB", "DDimer",
]
CYTOKINE_IDXS = list(range(29, 41))
CYTOKINE_NAMES = {
    29: "IL-5", 30: "IFN-alpha", 31: "IL-2", 32: "IL-6",
    33: "IL-1beta", 34: "IL-10", 35: "IFN-gamma", 36: "IL-8",
    37: "IL-17", 38: "IL-4", 39: "IL-12P70", 40: "TNF-alpha",
}
CATEGORICAL_VARS = [
    "Sex", "Cough", "Fever", "PleuralReaction", "PleuralEffusion",
    "PulmonaryConsolidation", "SMPP", "PE", "NP",
]
CATEGORICAL_IDXS = {
    41: "Bronchoscopy", 42: "ICU", 43: "Oxygen",
}

# ---------------------------------------------------------------------------
# 3. Helper functions
# ---------------------------------------------------------------------------
def sig_stars(p):
    if pd.isna(p):
        return ""
    if p < 0.001:
        return "***"
    elif p < 0.01:
        return "**"
    elif p < 0.05:
        return "*"
    return ""

def fmt_p(p):
    if pd.isna(p):
        return "N/A"
    if p < 0.001:
        return "<0.001"
    return f"{p:.3f}"

def fmt_mean_sd(m, s):
    return f"{m:.2f} +/- {s:.2f}"

def fmt_median_iqr(m, q1, q3):
    return f"{m:.2f} ({q1:.2f}-{q3:.2f})"

# ---------------------------------------------------------------------------
# 4. Analysis
# ---------------------------------------------------------------------------
results = []

# --- 4a. NORMAL variables (t-test, mean +/- SD, Levene) ---
for var in NORMAL_VARS:
    g0 = df_non[var].dropna().astype(float)
    g1 = df_rmpp[var].dropna().astype(float)
    mean0, sd0 = g0.mean(), g0.std(ddof=1)
    mean1, sd1 = g1.mean(), g1.std(ddof=1)
    _, levene_p = stats.levene(g0, g1)
    if levene_p >= 0.05:
        _, t_p = stats.ttest_ind(g0, g1, equal_var=True)
        tname = "Student t"
    else:
        _, t_p = stats.ttest_ind(g0, g1, equal_var=False)
        tname = "Welch t"
    results.append({
        "Variable": var,
        "nonRMPP": fmt_mean_sd(mean0, sd0),
        "RMPP":   fmt_mean_sd(mean1, sd1),
        "Test": tname,
        "P": t_p, "P_fmt": fmt_p(t_p), "Sig": sig_stars(t_p),
        "Levene_P": f"{levene_p:.4f}",
    })

# --- 4b. NON-NORMAL variables (Mann-Whitney, median IQR) ---
for var in NON_NORMAL_VARS:
    g0 = df_non[var].dropna().astype(float)
    g1 = df_rmpp[var].dropna().astype(float)
    med0, q1_0, q3_0 = np.percentile(g0, [50, 25, 75])
    med1, q1_1, q3_1 = np.percentile(g1, [50, 25, 75])
    try:
        _, mw_p = stats.mannwhitneyu(g0, g1, alternative="two-sided")
    except ValueError:
        mw_p = np.nan
    results.append({
        "Variable": var,
        "nonRMPP": fmt_median_iqr(med0, q1_0, q3_0),
        "RMPP":   fmt_median_iqr(med1, q1_1, q3_1),
        "Test": "Mann-Whitney",
        "P": mw_p, "P_fmt": fmt_p(mw_p), "Sig": sig_stars(mw_p),
        "Levene_P": "",
    })

# --- 4c. CYTOKINES (non-normal, Chinese-named columns) ---
for idx in CYTOKINE_IDXS:
    col = df.columns[idx]
    dname = CYTOKINE_NAMES.get(idx, col)
    g0 = df_non[col].dropna().astype(float)
    g1 = df_rmpp[col].dropna().astype(float)
    if len(g0) == 0 or len(g1) == 0:
        results.append({
            "Variable": dname, "nonRMPP": "--", "RMPP": "--",
            "Test": "Mann-Whitney", "P": np.nan, "P_fmt": "N/A",
            "Sig": "", "Levene_P": "",
        })
        continue
    med0, q1_0, q3_0 = np.percentile(g0, [50, 25, 75])
    med1, q1_1, q3_1 = np.percentile(g1, [50, 25, 75])
    try:
        _, mw_p = stats.mannwhitneyu(g0, g1, alternative="two-sided")
    except ValueError:
        mw_p = np.nan
    results.append({
        "Variable": dname,
        "nonRMPP": fmt_median_iqr(med0, q1_0, q3_0),
        "RMPP":   fmt_median_iqr(med1, q1_1, q3_1),
        "Test": "Mann-Whitney",
        "P": mw_p, "P_fmt": fmt_p(mw_p), "Sig": sig_stars(mw_p),
        "Levene_P": "",
    })

# --- 4d. CATEGORICAL variables (English-named columns) ---
for var in CATEGORICAL_VARS:
    s_non = df_non[var].dropna()
    s_rmpp = df_rmpp[var].dropna()
    valid = sorted(set(s_non.unique()) | set(s_rmpp.unique()))
    valid = [v for v in valid if v < 9]
    if not valid:
        results.append({
            "Variable": var, "nonRMPP": "No valid data", "RMPP": "No valid data",
            "Test": "--", "P": np.nan, "P_fmt": "--", "Sig": "", "Levene_P": "",
        })
        continue
    non_strs, rmpp_strs = [], []
    table_rows = []
    for lv in valid:
        n0 = int((s_non == lv).sum())
        n1 = int((s_rmpp == lv).sum())
        p0 = n0 / n_non * 100
        p1 = n1 / n_rmpp * 100
        non_strs.append(f"{int(lv)}={n0} ({p0:.1f}%)")
        rmpp_strs.append(f"{int(lv)}={n1} ({p1:.1f}%)")
        table_rows.append([n0, n1])
    tbl = np.array(table_rows, dtype=int)
    try:
        chi2, chi_p, dof, expected = stats.chi2_contingency(tbl)
        if np.any(expected < 5):
            if tbl.shape == (2, 2):
                _, fish_p = stats.fisher_exact(tbl)
                tlabel, pval = "Fisher", fish_p
            else:
                tlabel, pval = "Chi2*", chi_p
        else:
            tlabel, pval = "Chi2", chi_p
    except Exception:
        tlabel, pval = "--", np.nan
    results.append({
        "Variable": var,
        "nonRMPP": "; ".join(non_strs),
        "RMPP":   "; ".join(rmpp_strs),
        "Test": tlabel,
        "P": pval, "P_fmt": fmt_p(pval), "Sig": sig_stars(pval),
        "Levene_P": "",
    })

# --- 4e. CATEGORICAL variables (Chinese-named columns) ---
for idx, dname in CATEGORICAL_IDXS.items():
    col = df.columns[idx]
    s_non = df_non[col].dropna()
    s_rmpp = df_rmpp[col].dropna()
    valid = sorted(set(s_non.unique()) | set(s_rmpp.unique()))
    valid = [v for v in valid if v < 9]
    if not valid:
        results.append({
            "Variable": dname, "nonRMPP": "No valid data", "RMPP": "No valid data",
            "Test": "--", "P": np.nan, "P_fmt": "--", "Sig": "", "Levene_P": "",
        })
        continue
    non_strs, rmpp_strs = [], []
    table_rows = []
    for lv in valid:
        n0 = int((s_non == lv).sum())
        n1 = int((s_rmpp == lv).sum())
        p0 = n0 / n_non * 100
        p1 = n1 / n_rmpp * 100
        non_strs.append(f"{int(lv)}={n0} ({p0:.1f}%)")
        rmpp_strs.append(f"{int(lv)}={n1} ({p1:.1f}%)")
        table_rows.append([n0, n1])
    tbl = np.array(table_rows, dtype=int)
    try:
        chi2, chi_p, dof, expected = stats.chi2_contingency(tbl)
        if np.any(expected < 5):
            if tbl.shape == (2, 2):
                _, fish_p = stats.fisher_exact(tbl)
                tlabel, pval = "Fisher", fish_p
            else:
                tlabel, pval = "Chi2*", chi_p
        else:
            tlabel, pval = "Chi2", chi_p
    except Exception:
        tlabel, pval = "--", np.nan
    results.append({
        "Variable": dname,
        "nonRMPP": "; ".join(non_strs),
        "RMPP":   "; ".join(rmpp_strs),
        "Test": tlabel,
        "P": pval, "P_fmt": fmt_p(pval), "Sig": sig_stars(pval),
        "Levene_P": "",
    })

# ---------------------------------------------------------------------------
# 5. Console output
# ---------------------------------------------------------------------------
cols = ["Variable", "nonRMPP", "RMPP", "Test", "P_fmt", "Sig", "Levene_P"]
out_df = pd.DataFrame(results, columns=cols)

display_df = out_df.rename(columns={
    f"nonRMPP(n={n_non})": f"nonRMPP(n={n_non})",
    f"RMPP(n={n_rmpp})":   f"RMPP(n={n_rmpp})",
    "P_fmt": "P",
})

print()
print("=" * 130)
print(f"Table 1. Baseline Characteristics  |  nonRMPP (n={n_non}) vs RMPP (n={n_rmpp})")
print("=" * 130)
print(out_df[["Variable", "nonRMPP", "RMPP", "Test", "P_fmt", "Sig", "Levene_P"]].to_string(index=False, header=[
    "Variable", f"nonRMPP(n={n_non})", f"RMPP(n={n_rmpp})", "Test", "P", "Sig", "Levene_P"
]))
print("=" * 130)
print("Significance:  * p<0.05   ** p<0.01   *** p<0.001")
print("=" * 130)

# ---------------------------------------------------------------------------
# 6. Excel output
# ---------------------------------------------------------------------------
print(f"\n[INFO] Saving Excel to: {OUT_PATH}")

wb = Workbook()
ws = wb.active
ws.title = "Table1_Baseline"

hfont = Font(name="Arial", bold=True, size=11, color="FFFFFF")
hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
cfont = Font(name="Arial", size=10)
bdr = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
calign = Alignment(horizontal="center", vertical="center", wrap_text=True)
lalign = Alignment(horizontal="left", vertical="center", wrap_text=True)
sfill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

# Title
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
tc = ws.cell(row=1, column=1, value=f"Table 1. Baseline Characteristics - nonRMPP (n={n_non}) vs RMPP (n={n_rmpp})")
tc.font = Font(name="Arial", bold=True, size=14)
tc.alignment = Alignment(horizontal="center", vertical="center")

# Header
excel_headers = ["Variable", f"nonRMPP(n={n_non})", f"RMPP(n={n_rmpp})", "Test", "P", "Sig", "Levene_P"]
for ci, h in enumerate(excel_headers, 1):
    c = ws.cell(row=3, column=ci, value=h)
    c.font = hfont; c.fill = hfill; c.alignment = calign; c.border = bdr

# Data
for ri, (_, row) in enumerate(out_df.iterrows(), 4):
    vals = [
        row["Variable"],
        row["nonRMPP"],
        row["RMPP"],
        row["Test"],
        row["P_fmt"],
        row["Sig"],
        row["Levene_P"],
    ]
    for ci, val in enumerate(vals, 1):
        if pd.isna(val) or val == "":
            val = "--"
        c = ws.cell(row=ri, column=ci, value=val)
        c.font = cfont
        c.alignment = lalign if ci == 1 else calign
        c.border = bdr
    # Highlight significant rows
    sig = row["Sig"]
    if sig and sig not in ("", "--"):
        for ci2 in range(1, 8):
            ws.cell(row=ri, column=ci2).fill = sfill

# Column widths
for i, w in enumerate([22, 32, 32, 14, 12, 8, 12], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A4"

# Footer
fr = len(out_df) + 5
ws.merge_cells(start_row=fr, start_column=1, end_row=fr, end_column=7)
fc = ws.cell(row=fr, column=1, value="Significance: * p<0.05   ** p<0.01   *** p<0.001")
fc.font = Font(name="Arial", italic=True, size=9, color="666666")

wb.save(OUT_PATH)
print(f"[INFO] Done. Results saved to {OUT_PATH}")
