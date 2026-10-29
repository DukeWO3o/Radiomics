from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
from pathlib import Path

import joblib
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sklearn.metrics import roc_auc_score
from statsmodels.nonparametric.smoothers_lowess import lowess


SCRIPT_DIR = Path(__file__).resolve().parent
STRICT_ROOT = SCRIPT_DIR.parent
PIPELINE_ROOT = STRICT_ROOT.parent
sys.path.insert(0, str(PIPELINE_ROOT))
from pipeline_paths import (
    CALIBRATION_OUTPUT_ROOT,
    COMBINED_OUTPUT_ROOT,
    DEFAULT_MODEL_DATA_931,
    FIGURE_ROOT,
    OUTPUT_ROOT,
    SUPPLEMENTARY_OUTPUT_ROOT,
)

MODEL_DIR = COMBINED_OUTPUT_ROOT
DATA_FILE = DEFAULT_MODEL_DATA_931
PICKLE_HELPER = (
    PIPELINE_ROOT
    / "03_FeatureSelection_Step"
    / "strict_cv_pickle_helpers.py"
)
OR4_FILE = (
    SUPPLEMENTARY_OUTPUT_ROOT / "Online_Resource_4_model_reproducibility.xlsx"
)
FIGURE_DIR = FIGURE_ROOT
FIG6_STEM = "Fig6_Revised_SVM_Calibration_DCA"
BACKUP_DIR = OUTPUT_ROOT / "Backups" / "ThresholdCalibration"
JSON_OUTPUT = CALIBRATION_OUTPUT_ROOT / "Threshold_Calibration_Summary.json"
CALIBRATION_PANEL = CALIBRATION_OUTPUT_ROOT / "09_Calibration_SVM_with_95CI.svg"
POPPLER = shutil.which("pdftoppm")
THRESHOLDS = [0.4, 0.5, 0.6, 0.7]
PRIMARY_THRESHOLD = 0.5
N_BOOTSTRAP = 2000
RANDOM_STATE = 42


def load_pickle_helper():
    spec = importlib.util.spec_from_file_location(
        "_threshold_calibration_pickle_helper",
        PICKLE_HELPER,
    )
    if spec is None or spec.loader is None:
        raise ImportError(f"Cannot load helper module: {PICKLE_HELPER}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    setattr(
        sys.modules["__main__"],
        "StrictPreprocessSelector",
        module.StrictPreprocessSelector,
    )


def safe_logit(probabilities):
    probabilities = np.clip(probabilities, 1e-6, 1 - 1e-6)
    return np.log(probabilities / (1 - probabilities))


def calibration_parameters(labels, probabilities):
    logit = safe_logit(probabilities)
    slope_model = sm.GLM(
        labels,
        sm.add_constant(logit),
        family=sm.families.Binomial(),
    ).fit()
    intercept_model = sm.GLM(
        labels,
        np.ones((len(labels), 1)),
        family=sm.families.Binomial(),
        offset=logit,
    ).fit()
    observed = float(np.sum(labels))
    expected = float(np.sum(probabilities))
    return {
        "calibration_intercept": float(intercept_model.params[0]),
        "calibration_slope": float(slope_model.params[1]),
        "observed_events": observed,
        "expected_events": expected,
        "observed_to_expected_ratio": observed / expected,
    }


def threshold_metrics(labels, probabilities, threshold):
    predicted = probabilities >= threshold
    positive = labels == 1
    negative = ~positive
    tp = int(np.sum(predicted & positive))
    fp = int(np.sum(predicted & negative))
    fn = int(np.sum(~predicted & positive))
    tn = int(np.sum(~predicted & negative))
    sensitivity = tp / (tp + fn)
    specificity = tn / (tn + fp)
    ppv = tp / (tp + fp) if tp + fp else np.nan
    npv = tn / (tn + fn) if tn + fn else np.nan
    lr_positive = (
        sensitivity / (1 - specificity) if specificity < 1 else np.inf
    )
    lr_negative = (
        (1 - sensitivity) / specificity if specificity > 0 else np.inf
    )
    return {
        "threshold": float(threshold),
        "tp": tp,
        "fp": fp,
        "fn": fn,
        "tn": tn,
        "sensitivity": sensitivity,
        "specificity": specificity,
        "ppv": ppv,
        "npv": npv,
        "lr_positive": lr_positive,
        "lr_negative": lr_negative,
    }


def percentile_ci(values):
    finite = np.asarray(values, dtype=float)
    finite = finite[np.isfinite(finite)]
    if not len(finite):
        return [None, None]
    return [
        float(np.percentile(finite, 2.5)),
        float(np.percentile(finite, 97.5)),
    ]


def bootstrap_analysis(labels, probabilities):
    rng = np.random.RandomState(RANDOM_STATE)
    metric_names = [
        "sensitivity",
        "specificity",
        "ppv",
        "npv",
        "lr_positive",
        "lr_negative",
    ]
    metric_values = {
        f"{metric}@{threshold}": []
        for threshold in THRESHOLDS
        for metric in metric_names
    }
    calibration_values = {
        "calibration_intercept": [],
        "calibration_slope": [],
        "observed_to_expected_ratio": [],
    }
    for _ in range(N_BOOTSTRAP):
        indices = rng.choice(len(labels), len(labels), replace=True)
        sampled_labels = labels[indices]
        if len(np.unique(sampled_labels)) < 2:
            continue
        sampled_probabilities = probabilities[indices]
        for threshold in THRESHOLDS:
            metrics = threshold_metrics(
                sampled_labels,
                sampled_probabilities,
                threshold,
            )
            for metric in metric_names:
                metric_values[f"{metric}@{threshold}"].append(
                    metrics[metric]
                )
        calibration = calibration_parameters(
            sampled_labels,
            sampled_probabilities,
        )
        for metric in calibration_values:
            calibration_values[metric].append(calibration[metric])
    return {
        metric: percentile_ci(values)
        for metric, values in {**metric_values, **calibration_values}.items()
    }


def build_summary(labels, probabilities):
    point_thresholds = [
        threshold_metrics(labels, probabilities, threshold)
        for threshold in THRESHOLDS
    ]
    intervals = bootstrap_analysis(labels, probabilities)
    summary = {
        "n": int(len(labels)),
        "events": int(np.sum(labels)),
        "nonevents": int(np.sum(labels == 0)),
        "prevalence": float(np.mean(labels)),
        "auc": float(roc_auc_score(labels, probabilities)),
        "mean_predicted_probability": float(np.mean(probabilities)),
        "primary_threshold": PRIMARY_THRESHOLD,
        "primary_threshold_status": "fixed a priori; not optimized on validation",
        "svm_class_weight": "balanced",
        "bootstrap_resamples": N_BOOTSTRAP,
        "random_state": RANDOM_STATE,
        "thresholds": point_thresholds,
        "bootstrap_95ci": intervals,
        "calibration": calibration_parameters(labels, probabilities),
    }
    return summary


def calibration_band(labels, probabilities, grid):
    rng = np.random.RandomState(RANDOM_STATE)
    curves = []
    for _ in range(1000):
        indices = rng.choice(len(labels), len(labels), replace=True)
        sampled_labels = labels[indices]
        if len(np.unique(sampled_labels)) < 2:
            continue
        sampled_probabilities = probabilities[indices]
        curve = lowess(
            sampled_labels,
            sampled_probabilities,
            xvals=grid,
            frac=0.8,
            it=3,
            return_sorted=False,
        )
        curves.append(np.clip(curve, 0, 1))
    curves = np.asarray(curves)
    return (
        np.percentile(curves, 2.5, axis=0),
        np.percentile(curves, 97.5, axis=0),
    )


def dca_curves(labels, probabilities):
    thresholds = np.arange(0, 1.01, 0.01)
    prevalence = float(np.mean(labels))
    model = []
    treat_all = []
    for threshold in thresholds:
        if threshold >= 1.0:
            model.append(0.0)
            treat_all.append(0.0)
            continue
        predicted = probabilities >= threshold
        tp = np.sum((labels == 1) & predicted)
        fp = np.sum((labels == 0) & predicted)
        weight = threshold / (1 - threshold)
        model.append(tp / len(labels) - fp / len(labels) * weight)
        treat_all.append(prevalence - (1 - prevalence) * weight)
    return thresholds, np.asarray(model), np.asarray(treat_all)


def format_estimate(value):
    if value is None or not np.isfinite(value):
        return "NA"
    return f"{value:.3f}"


def format_ci(metric, interval):
    low, high = interval
    if low is None or high is None:
        return "NA"
    return f"{format_estimate(low)} to {format_estimate(high)}"


def update_or4(summary):
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    backup = BACKUP_DIR / OR4_FILE.name
    if not backup.exists():
        shutil.copy2(OR4_FILE, backup)

    workbook = load_workbook(OR4_FILE)
    sheet_name = "Threshold_Calibration_SVM"
    if sheet_name in workbook.sheetnames:
        workbook.remove(workbook[sheet_name])
    sheet = workbook.create_sheet(
        sheet_name,
        index=workbook.sheetnames.index("Model_Performance") + 1,
    )
    sheet.sheet_view.showGridLines = False

    title_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    pale_fill = PatternFill("solid", fgColor="F7F9FA")
    thin_gray = Side(style="thin", color="D9D9D9")
    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )

    sheet.merge_cells("A1:J1")
    sheet["A1"] = "SVM threshold performance and calibration"
    sheet["A1"].fill = title_fill
    sheet["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 26

    sheet.merge_cells("A2:J2")
    sheet["A2"] = (
        f"Locked internal validation set: n={summary['n']}, "
        f"RMPP events={summary['events']}, prevalence="
        f"{summary['prevalence']:.1%}. Bootstrap resamples="
        f"{summary['bootstrap_resamples']}; random seed={summary['random_state']}. "
        "The 0.5 threshold was fixed a priori and was not optimized in the "
        "validation set. The SVM used class_weight='balanced'."
    )
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[2].height = 48

    sheet.merge_cells("A4:J4")
    sheet["A4"] = "A. Threshold-specific performance"
    sheet["A4"].fill = section_fill
    sheet["A4"].font = Font(bold=True, size=11, color="17242A")

    headers = [
        "Threshold",
        "TP / FP / FN / TN",
        "Sensitivity (95% CI)",
        "Specificity (95% CI)",
        "PPV (95% CI)",
        "NPV (95% CI)",
        "LR+ (95% CI)",
        "LR- (95% CI)",
        "Role",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=column, value=header)
        cell.fill = title_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border
    sheet.row_dimensions[5].height = 36

    intervals = summary["bootstrap_95ci"]
    for row_index, metrics in enumerate(summary["thresholds"], start=6):
        threshold = metrics["threshold"]
        values = [
            f"{threshold:.1f}",
            f"{metrics['tp']} / {metrics['fp']} / {metrics['fn']} / {metrics['tn']}",
            f"{metrics['sensitivity']:.3f} "
            f"({format_ci('sensitivity', intervals[f'sensitivity@{threshold:.1f}'])})",
            f"{metrics['specificity']:.3f} "
            f"({format_ci('specificity', intervals[f'specificity@{threshold:.1f}'])})",
            f"{metrics['ppv']:.3f} "
            f"({format_ci('ppv', intervals[f'ppv@{threshold:.1f}'])})",
            f"{metrics['npv']:.3f} "
            f"({format_ci('npv', intervals[f'npv@{threshold:.1f}'])})",
            f"{metrics['lr_positive']:.3f} "
            f"({format_ci('lr_positive', intervals[f'lr_positive@{threshold:.1f}'])})",
            f"{metrics['lr_negative']:.3f} "
            f"({format_ci('lr_negative', intervals[f'lr_negative@{threshold:.1f}'])})",
            "Primary" if threshold == PRIMARY_THRESHOLD else "Descriptive",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if column == 9 else "center",
                vertical="center",
                wrap_text=True,
            )
            if row_index % 2 == 0:
                cell.fill = pale_fill

    sheet.merge_cells("A11:J11")
    sheet["A11"] = "B. Calibration and average prediction error"
    sheet["A11"].fill = section_fill
    sheet["A11"].font = Font(bold=True, size=11, color="17242A")

    calibration_headers = ["Metric", "Estimate", "95% CI", "Interpretation"]
    for column, header in enumerate(calibration_headers, start=1):
        cell = sheet.cell(row=12, column=column, value=header)
        cell.fill = title_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border
    calibration = summary["calibration"]
    calibration_rows = [
        (
            "Calibration intercept",
            f"{calibration['calibration_intercept']:.3f}",
            format_ci(
                "calibration_intercept",
                intervals["calibration_intercept"],
            ),
            "Perfect calibration has intercept 0.",
        ),
        (
            "Calibration slope",
            f"{calibration['calibration_slope']:.3f}",
            format_ci("calibration_slope", intervals["calibration_slope"]),
            "Perfect calibration has slope 1.",
        ),
        (
            "Observed-to-expected ratio",
            f"{calibration['observed_to_expected_ratio']:.3f}",
            format_ci(
                "observed_to_expected_ratio",
                intervals["observed_to_expected_ratio"],
            ),
            "Perfect average calibration has ratio 1.",
        ),
        (
            "Brier score",
            "0.200",
            "Not applicable",
            "Overall probabilistic prediction error.",
        ),
    ]
    for row_index, values in enumerate(calibration_rows, start=13):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if column == 4 else "center",
                vertical="center",
                wrap_text=True,
            )

    sheet.merge_cells("A18:J18")
    sheet["A18"] = (
        "Bootstrap intervals resample patients and do not represent external "
        "validation. Balanced class weights can alter absolute probabilities. "
        "Because observed prevalence was 59.9%, PPV and calibration may be "
        "lower in lower-prevalence populations; recalibration should be "
        "assessed externally before transportability is claimed."
    )
    sheet["A18"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[18].height = 54

    widths = {
        "A": 13,
        "B": 22,
        "C": 22,
        "D": 22,
        "E": 22,
        "F": 22,
        "G": 22,
        "H": 22,
        "I": 14,
        "J": 4,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = "A5:I9"
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    readme = workbook["README"]
    for row_index in range(readme.max_row, 1, -1):
        if (
            readme.cell(row_index, 1).value
            == "Threshold_Calibration_SVM"
        ):
            readme.delete_rows(row_index)
    next_row = readme.max_row + 1
    readme.cell(
        row=next_row,
        column=1,
        value="Threshold_Calibration_SVM",
    )
    readme.cell(
        row=next_row,
        column=2,
        value=(
            "Prespecified 0.5 threshold performance, descriptive "
            "0.4/0.6/0.7 analyses, and bootstrap calibration intervals."
        ),
    )
    readme.column_dimensions["A"].width = 34
    readme.column_dimensions["B"].width = 110
    for row in readme.iter_rows(min_row=1, max_row=next_row, max_col=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
    for row in readme.iter_rows():
        if row[0].value == "Synchronization date":
            row[1].value = "2026-10-25"

    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    temporary = OR4_FILE.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(OR4_FILE)


def render_pdf_to_png(pdf_path, png_path, dpi=300):
    if not POPPLER:
        raise FileNotFoundError(f"Poppler pdftoppm not found: {POPPLER}")
    if png_path.exists():
        png_path.unlink()
    output_stem = png_path.with_suffix("")
    subprocess.run(
        [
            str(POPPLER),
            "-png",
            "-r",
            str(dpi),
            "-singlefile",
            str(pdf_path),
            str(output_stem),
        ],
        check=True,
        capture_output=True,
    )
    if not png_path.exists():
        raise FileNotFoundError(f"PNG was not created: {png_path}")


def build_fig6(summary, labels, probabilities):
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    for extension in ("png", "pdf", "svg"):
        source = FIGURE_DIR / f"{FIG6_STEM}.{extension}"
        backup = BACKUP_DIR / source.name
        if source.exists() and not backup.exists():
            shutil.copy2(source, backup)

    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "DejaVu Sans", "sans-serif"],
            "font.size": 11,
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
        }
    )
    curve_grid = np.linspace(0.0, 1.0, 201)
    band_grid = np.linspace(0.40, 0.80, 120)
    observed_curve = np.clip(
        lowess(
            labels,
            probabilities,
            xvals=curve_grid,
            frac=0.8,
            it=3,
            return_sorted=False,
        ),
        0,
        1,
    )
    lower, upper = calibration_band(labels, probabilities, band_grid)
    thresholds, model_nb, treat_all = dca_curves(labels, probabilities)

    figure, axes = plt.subplots(
        1,
        2,
        figsize=(14, 7.5),
        facecolor="white",
    )
    calibration_axis, dca_axis = axes
    calibration_axis.fill_between(
        band_grid,
        lower,
        upper,
        color="#9ECAE1",
        alpha=0.45,
        label="Bootstrap 95% CI",
    )
    calibration_axis.plot(
        curve_grid,
        observed_curve,
        color="#111111",
        linewidth=2.2,
        label="SVM calibration",
    )
    calibration_axis.plot(
        [0, 1],
        [0, 1],
        color="#777777",
        linestyle="--",
        linewidth=1.4,
        label="Ideal",
    )
    calibration_axis.set_xlim(0, 1)
    calibration_axis.set_ylim(0, 1)
    calibration_axis.set_xlabel("Predicted probability", fontsize=13)
    calibration_axis.set_ylabel("Observed probability", fontsize=13)
    calibration_axis.set_title(
        "Calibration curve - SVM (Brier = 0.200)",
        fontsize=15,
    )
    calibration_axis.text(
        0.04,
        0.96,
        (
            f"Intercept {summary['calibration']['calibration_intercept']:.2f} "
            f"(95% CI {format_ci('calibration_intercept', summary['bootstrap_95ci']['calibration_intercept'])})\n"
            f"Slope {summary['calibration']['calibration_slope']:.2f} "
            f"(95% CI {format_ci('calibration_slope', summary['bootstrap_95ci']['calibration_slope'])})\n"
            f"O:E {summary['calibration']['observed_to_expected_ratio']:.2f} "
            f"(95% CI {format_ci('observed_to_expected_ratio', summary['bootstrap_95ci']['observed_to_expected_ratio'])})"
        ),
        transform=calibration_axis.transAxes,
        ha="left",
        va="top",
        fontsize=9.5,
        bbox={"facecolor": "white", "edgecolor": "#D9D9D9", "pad": 4},
    )
    calibration_axis.legend(
        loc="lower right",
        frameon=False,
        fontsize=10,
    )
    calibration_axis.grid(
        True,
        color="#E5E5E5",
        linewidth=0.7,
    )
    calibration_axis.set_axisbelow(True)
    calibration_axis.spines["top"].set_visible(False)
    calibration_axis.spines["right"].set_visible(False)
    calibration_axis.text(
        -0.10,
        1.02,
        "A",
        transform=calibration_axis.transAxes,
        fontsize=16,
        fontweight="bold",
        ha="left",
        va="bottom",
    )

    dca_axis.plot(
        thresholds,
        model_nb,
        color="#C62828",
        linewidth=2.2,
        label="SVM",
    )
    dca_axis.plot(
        thresholds,
        treat_all,
        color="#777777",
        linestyle="--",
        linewidth=1.4,
        label="Treat all",
    )
    dca_axis.plot(
        thresholds,
        np.zeros_like(thresholds),
        color="#202020",
        linewidth=1.2,
        label="Treat none",
    )
    dca_axis.set_xlim(0, 1)
    dca_axis.set_ylim(-0.05, 0.65)
    dca_axis.set_xlabel("Threshold probability", fontsize=13)
    dca_axis.set_ylabel("Net benefit", fontsize=13)
    dca_axis.set_title("Decision curve analysis - SVM", fontsize=15)
    dca_axis.legend(loc="upper right", frameon=False, fontsize=10)
    dca_axis.grid(True, color="#E5E5E5", linewidth=0.7)
    dca_axis.set_axisbelow(True)
    dca_axis.spines["top"].set_visible(False)
    dca_axis.spines["right"].set_visible(False)
    dca_axis.text(
        -0.10,
        1.02,
        "B",
        transform=dca_axis.transAxes,
        fontsize=16,
        fontweight="bold",
        ha="left",
        va="bottom",
    )
    figure.tight_layout()

    panel_figure, panel_axis = plt.subplots(figsize=(8, 8), facecolor="white")
    panel_axis.fill_between(
        band_grid,
        lower,
        upper,
        color="#9ECAE1",
        alpha=0.45,
        label="Bootstrap 95% CI",
    )
    panel_axis.plot(
        curve_grid,
        observed_curve,
        color="#111111",
        linewidth=2.2,
        label="SVM calibration",
    )
    panel_axis.plot(
        [0, 1],
        [0, 1],
        color="#777777",
        linestyle="--",
        label="Ideal",
    )
    panel_axis.set_xlim(0, 1)
    panel_axis.set_ylim(0, 1)
    panel_axis.set_xlabel("Predicted probability", fontsize=16)
    panel_axis.set_ylabel("Observed probability", fontsize=16)
    panel_axis.set_title("Calibration curve - SVM (Brier = 0.200)", fontsize=18)
    panel_axis.legend(loc="lower right", frameon=False)
    panel_axis.grid(True, color="#E5E5E5", linewidth=0.7)
    panel_figure.savefig(
        CALIBRATION_PANEL,
        format="svg",
        dpi=300,
        bbox_inches="tight",
        facecolor="white",
    )
    plt.close(panel_figure)

    svg_path = FIGURE_DIR / f"{FIG6_STEM}.svg"
    png_path = FIGURE_DIR / f"{FIG6_STEM}.png"
    pdf_path = FIGURE_DIR / f"{FIG6_STEM}.pdf"
    figure.savefig(
        svg_path,
        format="svg",
        dpi=300,
        bbox_inches="tight",
        facecolor="white",
    )
    figure.savefig(
        pdf_path,
        format="pdf",
        dpi=300,
        bbox_inches="tight",
        facecolor="white",
    )
    plt.close(figure)
    render_pdf_to_png(pdf_path, png_path, dpi=300)


def main():
    CALIBRATION_OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    load_pickle_helper()
    data = pd.read_excel(DATA_FILE).reset_index(drop=True)
    manifest = joblib.load(
        MODEL_DIR / "00_Validation_Patient_Manifest.pkl"
    )
    model = joblib.load(MODEL_DIR / "model_SVM.pkl")
    merged = manifest.merge(
        data,
        on=["image", "label"],
        how="left",
        validate="one_to_one",
    )
    labels = merged["label"].to_numpy()
    probabilities = model.predict_proba(merged)[:, 1]

    summary = build_summary(labels, probabilities)
    JSON_OUTPUT.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    update_or4(summary)
    build_fig6(summary, labels, probabilities)
    print(JSON_OUTPUT)
    print(OR4_FILE)
    print(FIGURE_DIR / f"{FIG6_STEM}.png")
    print(FIGURE_DIR / f"{FIG6_STEM}.svg")
    print(FIGURE_DIR / f"{FIG6_STEM}.pdf")


if __name__ == "__main__":
    main()
