from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


PIPELINE_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PIPELINE_ROOT))
from pipeline_paths import (
    PIXEL_COMBAT_ROOT,
    PIXEL_NO_COMBAT_ROOT,
    PIXEL_PREPROCESS_ROOT,
    PIXEL_SUMMARY_ROOT,
    SUPPLEMENTARY_OUTPUT_ROOT,
)

DATA_DIR = PIXEL_PREPROCESS_ROOT
NO_COMBAT_DIR = PIXEL_NO_COMBAT_ROOT
COMBAT_DIR = PIXEL_COMBAT_ROOT
OUTPUT = (
    SUPPLEMENTARY_OUTPUT_ROOT
    / "Online_Resource_11_PixelSpacing_geometric_sensitivity.xlsx"
)


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def main() -> None:
    manifest = pd.read_csv(
        DATA_DIR / "01_E_processing_manifest.csv",
        encoding="utf-8-sig",
    )
    icc = read_json(DATA_DIR / "03_ICC_summary_pixel_spacing_497.json")

    geometry = pd.DataFrame(
        [
            {
                "Item": "Cases with PixelSpacing",
                "Value": len(manifest),
                "Explanation": (
                    "The 497-case subset was selected only by DICOM "
                    "PixelSpacing availability and used as a sensitivity cohort."
                ),
            },
            {
                "Item": "Median effective pixel size in final 1024 canvas, x",
                "Value": float(manifest["final_spacing_x_mm"].median()),
                "Explanation": (
                    "Effective spacing after uniform aspect-preserving scaling; "
                    "not fixed at 0.5 mm."
                ),
            },
            {
                "Item": "IQR effective pixel size, x",
                "Value": (
                    f"{manifest['final_spacing_x_mm'].quantile(0.25):.4f}-"
                    f"{manifest['final_spacing_x_mm'].quantile(0.75):.4f}"
                ),
                "Explanation": "25th-75th percentile.",
            },
            {
                "Item": "Range effective pixel size, x",
                "Value": (
                    f"{manifest['final_spacing_x_mm'].min():.4f}-"
                    f"{manifest['final_spacing_x_mm'].max():.4f}"
                ),
                "Explanation": "Observed minimum-maximum.",
            },
            {
                "Item": "Median effective pixel size in final 1024 canvas, y",
                "Value": float(manifest["final_spacing_y_mm"].median()),
                "Explanation": (
                    "Effective spacing after uniform aspect-preserving scaling; "
                    "not fixed at 0.5 mm."
                ),
            },
            {
                "Item": "IQR effective pixel size, y",
                "Value": (
                    f"{manifest['final_spacing_y_mm'].quantile(0.25):.4f}-"
                    f"{manifest['final_spacing_y_mm'].quantile(0.75):.4f}"
                ),
                "Explanation": "25th-75th percentile.",
            },
            {
                "Item": "Range effective pixel size, y",
                "Value": (
                    f"{manifest['final_spacing_y_mm'].min():.4f}-"
                    f"{manifest['final_spacing_y_mm'].max():.4f}"
                ),
                "Explanation": "Observed minimum-maximum.",
            },
            {
                "Item": "Field of view",
                "Value": "Original DICOM FOV",
                "Explanation": (
                    "Aspect-preserving scaling retained the anatomical field of "
                    "view. Zero padding added no anatomical information."
                ),
            },
            {
                "Item": "Relationship between 0.5 mm and 1024",
                "Value": "Sequential processing",
                "Explanation": (
                    "1024 x 1024 is a computational canvas, not a fixed "
                    "512 x 512-mm FOV. The final scaling changed effective pixel "
                    "size. PyRadiomics used 0.5 x 0.5 mm as the prespecified "
                    "feature-extraction target."
                ),
            },
        ]
    )

    cohort = pd.DataFrame(
        [
            ["Total", len(manifest)],
            ["RMPP", int((manifest["label"] == 1).sum())],
            ["non-RMPP", int((manifest["label"] == 0).sum())],
            ["Training", 397],
            ["Locked internal validation", 100],
        ],
        columns=["Characteristic", "N"],
    )
    batches = (
        manifest.groupby(["manufacturer", "manufacturer_model"])
        .size()
        .rename("N")
        .reset_index()
        .sort_values("N", ascending=False)
    )
    icc_frame = pd.DataFrame(
        [{"Metric": key, "Value": value} for key, value in icc.items()]
    )

    performance_rows = []
    best_rows = []
    parameter_rows = []
    for branch_name, branch_dir in [
        ("No ComBat", NO_COMBAT_DIR),
        ("ComBat", COMBAT_DIR),
    ]:
        model_dir = branch_dir / "\u8054\u5408\u6a21\u578b"
        performance = pd.read_excel(
            model_dir / "04_Model_Performance.xlsx",
            index_col=0,
        )
        config = read_json(model_dir / "model_config.json")
        best_model = config["best_model"]
        best = performance.loc[best_model]
        best_rows.append(
            {
                "Branch": branch_name,
                "Selected model": best_model,
                "Training CV AUC": float(best["Selection_CV_AUC"]),
                "Validation AUC": float(best["Valid_AUC"]),
                "95% CI low": float(best["Valid_AUC_95CI_Low"]),
                "95% CI high": float(best["Valid_AUC_95CI_High"]),
                "Recall": float(best["Valid_Recall"]),
                "Specificity": float(best["Valid_Specificity"]),
                "Brier": float(best["Valid_Brier"]),
            }
        )
        for model_name, values in performance.iterrows():
            performance_rows.append(
                {
                    "Branch": branch_name,
                    "Model": model_name,
                    "Training CV AUC": float(values["Selection_CV_AUC"]),
                    "Validation AUC": float(values["Valid_AUC"]),
                    "Validation AUC 95% CI low": float(
                        values["Valid_AUC_95CI_Low"]
                    ),
                    "Validation AUC 95% CI high": float(
                        values["Valid_AUC_95CI_High"]
                    ),
                    "Recall": float(values["Valid_Recall"]),
                    "Specificity": float(values["Valid_Specificity"]),
                    "Brier": float(values["Valid_Brier"]),
                }
            )
        fixed_parameters = read_json(
            branch_dir / "14_Tuning_Best_Parameters.json"
        )
        for model_name, parameters in fixed_parameters.items():
            parameter_rows.append(
                {
                    "Branch": branch_name,
                    "Model": model_name,
                    "Selected parameters": json.dumps(
                        parameters,
                        ensure_ascii=False,
                        sort_keys=True,
                    ),
                }
            )

    svm_comparison = pd.read_excel(
        PIXEL_SUMMARY_ROOT / "06_PixelSpacing497_Model_Summary.xlsx",
        sheet_name="Same_SVM_DeLong",
    )
    readme = pd.DataFrame(
        {
            "Item": [
                "Purpose",
                "Geometry clarification",
                "Sensitivity cohort",
                "Primary model status",
                "Interpretation",
            ],
            "Description": [
                (
                    "This resource supports the response to the reviewer request "
                    "on the relationship between 0.5 x 0.5-mm spacing and the "
                    "1024 x 1024 image matrix."
                ),
                (
                    "The 0.5 x 0.5-mm value was an intermediate resampling target "
                    "and the PyRadiomics feature-extraction spacing. Aspect-preserving "
                    "scaling and zero padding then produced a 1024 x 1024 computational "
                    "canvas. Final effective pixel size therefore varied, and the "
                    "matrix was not interpreted as a fixed 512 x 512-mm FOV."
                ),
                (
                    "The 497 children with available PixelSpacing were analysed "
                    "separately. Actual physical spacing was assigned before "
                    "radiomics resampling to 0.5 x 0.5 mm."
                ),
                (
                    "The primary manuscript model remains the original 931-case "
                    "SVM. The PixelSpacing analysis is exploratory and was not used "
                    "to replace the primary model."
                ),
                (
                    "The subset was defined by DICOM field availability, its "
                    "validation set contained 100 children, and no external centre "
                    "or unseen-device validation was available. Results should not "
                    "be interpreted as evidence of external generalisation."
                ),
            ],
        }
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        geometry.to_excel(writer, sheet_name="Geometry_Relationship", index=False)
        cohort.to_excel(writer, sheet_name="Cohort", index=False)
        batches.to_excel(writer, sheet_name="Batches", index=False)
        icc_frame.to_excel(writer, sheet_name="ICC", index=False)
        pd.DataFrame(best_rows).to_excel(
            writer,
            sheet_name="Best_Models",
            index=False,
        )
        pd.DataFrame(performance_rows).to_excel(
            writer,
            sheet_name="All_Model_Performance",
            index=False,
        )
        svm_comparison.to_excel(
            writer,
            sheet_name="Same_SVM_DeLong",
            index=False,
        )
        pd.DataFrame(parameter_rows).to_excel(
            writer,
            sheet_name="Best_Parameters",
            index=False,
        )

    workbook = load_workbook(OUTPUT)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = border
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )
                cell.border = border
        for column in sheet.columns:
            values = [
                str(cell.value) if cell.value is not None else ""
                for cell in column
            ]
            width = min(
                max(max((len(value) for value in values), default=8) + 2, 12),
                60,
            )
            sheet.column_dimensions[column[0].column_letter].width = width
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
