from __future__ import annotations

import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy.stats import spearmanr


PIPELINE_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PIPELINE_ROOT))
from pipeline_paths import (
    COMBINED_OUTPUT_ROOT,
    CORRELATION_OUTPUT_ROOT,
    DEFAULT_MODEL_DATA_931,
    SUPPLEMENTARY_OUTPUT_ROOT,
)

INPUT_FILE = DEFAULT_MODEL_DATA_931
MODEL_CONFIG = COMBINED_OUTPUT_ROOT / "model_config.json"
OUTPUT_DIR = CORRELATION_OUTPUT_ROOT
SUPPLEMENTARY_DIR = SUPPLEMENTARY_OUTPUT_ROOT

OUTPUT_XLSX = OUTPUT_DIR / "J_Correlation_Age_ROI.xlsx"
OUTPUT_PNG = OUTPUT_DIR / "J_Correlation_Age_ROI_heatmap.png"
ONLINE_XLSX = (
    SUPPLEMENTARY_DIR / "Online_Resource_7_final_feature_correlation.xlsx"
)
ONLINE_PNG = (
    SUPPLEMENTARY_DIR
    / "Online_Resource_7_final_feature_correlation_heatmap.png"
)

N_BOOTSTRAP = 1000
RANDOM_STATE = 42


def benjamini_hochberg(p_values: np.ndarray) -> np.ndarray:
    p_values = np.asarray(p_values, dtype=float)
    adjusted = np.full_like(p_values, np.nan, dtype=float)
    valid = ~np.isnan(p_values)
    p = p_values[valid]
    if p.size == 0:
        return adjusted
    order = np.argsort(p)
    ranked = p[order]
    m = p.size
    q = ranked * m / np.arange(1, m + 1)
    q = np.minimum.accumulate(q[::-1])[::-1]
    q = np.clip(q, 0, 1)
    restored = np.empty_like(q)
    restored[order] = q
    adjusted[valid] = restored
    return adjusted


def bootstrap_spearman_ci(
    x: np.ndarray,
    y: np.ndarray,
    n_bootstrap: int = N_BOOTSTRAP,
    random_state: int = RANDOM_STATE,
) -> tuple[float, float]:
    rng = np.random.default_rng(random_state)
    n = len(x)
    values = []
    for _ in range(n_bootstrap):
        idx = rng.integers(0, n, size=n)
        xb = x[idx]
        yb = y[idx]
        if np.unique(xb).size < 2 or np.unique(yb).size < 2:
            continue
        rho, _ = spearmanr(xb, yb, nan_policy="omit")
        if np.isfinite(rho):
            values.append(float(rho))
    if not values:
        return np.nan, np.nan
    low, high = np.percentile(values, [2.5, 97.5])
    return float(low), float(high)


def format_display_name(feature: str) -> str:
    return feature.replace("_", " ")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    SUPPLEMENTARY_DIR.mkdir(parents=True, exist_ok=True)

    with MODEL_CONFIG.open("r", encoding="utf-8") as stream:
        config = json.load(stream)

    protected = set(config["protected_features"])
    final_radiomics = [
        feature
        for feature in config["final_features"]
        if feature not in protected
    ]
    if len(final_radiomics) != 8:
        raise ValueError(
            f"Expected 8 final radiomics features, found {len(final_radiomics)}"
        )

    data = pd.read_excel(INPUT_FILE)
    required = ["age", "shape2D_PixelSurface", *final_radiomics]
    missing_columns = [column for column in required if column not in data.columns]
    if missing_columns:
        raise ValueError(f"Missing required columns: {missing_columns}")


    data = data.copy()
    data["ROI_area_pixels"] = pd.to_numeric(
        data["shape2D_PixelSurface"], errors="coerce"
    )
    correlation_rows = []
    targets = [
        ("Age", "age"),
        ("ROI area", "ROI_area_pixels"),
    ]

    for target_label, target_column in targets:
        for feature in final_radiomics:
            pair = data[[feature, target_column]].dropna()
            x = pair[feature].to_numpy(dtype=float)
            y = pair[target_column].to_numpy(dtype=float)
            rho, p_value = spearmanr(x, y)

            definitional = (
                target_column == "ROI_area_pixels"
                and feature == "shape2D_PixelSurface"
            )
            if definitional:
                ci_low = 1.0
                ci_high = 1.0
                p_adj = np.nan
                testable = False
            else:
                ci_low, ci_high = bootstrap_spearman_ci(x, y)
                p_adj = np.nan
                testable = True

            correlation_rows.append(
                {
                    "Feature": feature,
                    "Feature_display": format_display_name(feature),
                    "Target": target_label,
                    "N": len(pair),
                    "Spearman_rho": float(rho),
                    "CI95_low": ci_low,
                    "CI95_high": ci_high,
                    "P_raw": float(p_value),
                    "P_FDR": p_adj,
                    "Testable": testable,
                    "Interpretation": (
                        "Definitional identity: shape2D_PixelSurface is the "
                        "pixel-based ROI area"
                        if definitional
                        else ""
                    ),
                }
            )

    results = pd.DataFrame(correlation_rows)
    testable_mask = results["Testable"]
    results.loc[testable_mask, "P_FDR"] = benjamini_hochberg(
        results.loc[testable_mask, "P_raw"].to_numpy()
    )
    results["Significant_FDR_0.05"] = (
        results["P_FDR"].notna() & (results["P_FDR"] < 0.05)
    )

    results = results.sort_values(
        ["Target", "P_raw"], ascending=[True, True]
    ).reset_index(drop=True)

    metadata = pd.DataFrame(
        {
            "Item": [
                "Input file",
                "Final model configuration",
                "Analysis type",
                "Correlation method",
                "ROI area definition",
                "Physical ROI area conversion",
                "Confidence intervals",
                "Multiple testing",
                "Definitional self-correlation",
                "Final radiomics features",
            ],
            "Value": [
                str(INPUT_FILE),
                str(MODEL_CONFIG),
                "Post hoc secondary analysis after the final model was locked",
                "Spearman rank correlation",
                "shape2D_PixelSurface, corresponding to the number of positive "
                "pixels in the 1024 x 1024 binarized mask",
                "Not applicable; the 1024 x 1024 canvas is not a calibrated "
                "physical grid.",
                f"Percentile bootstrap, {N_BOOTSTRAP} resamples, random seed "
                f"{RANDOM_STATE}",
                "Benjamini-Hochberg FDR across 15 testable correlations "
                "(8 age correlations plus 7 non-definitional ROI-area correlations)",
                "shape2D_PixelSurface versus ROI area is a structural identity "
                "and was not included in inferential testing",
                "; ".join(final_radiomics),
            ],
        }
    )

    descriptive = data[
        ["age", "ROI_area_pixels", *final_radiomics]
    ].describe().T.reset_index().rename(columns={"index": "Variable"})

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        metadata.to_excel(writer, sheet_name="README", index=False)
        results.to_excel(writer, sheet_name="Correlation_Results", index=False)
        descriptive.to_excel(writer, sheet_name="Descriptive_Statistics", index=False)

    workbook = load_workbook(OUTPUT_XLSX)
    header_fill = PatternFill("solid", fgColor="17324D")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for row in sheet.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(
                    vertical="top", wrap_text=True
                )
        for column_cells in sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            max_len = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            sheet.column_dimensions[letter].width = min(max(12, max_len + 2), 60)

    results_sheet = workbook["Correlation_Results"]
    headers = {
        cell.value: cell.column
        for cell in results_sheet[1]
    }
    for row_index in range(2, results_sheet.max_row + 1):
        significance = results_sheet.cell(
            row_index, headers["Significant_FDR_0.05"]
        ).value
        if significance:
            for column in range(1, results_sheet.max_column + 1):
                results_sheet.cell(row_index, column).fill = PatternFill(
                    "solid", fgColor="E2EFDA"
                )

    workbook.save(OUTPUT_XLSX)

    heatmap_data = (
        results.pivot(index="Feature_display", columns="Target", values="Spearman_rho")
        .reindex([format_display_name(x) for x in final_radiomics])
    )
    annotation = heatmap_data.copy().astype(object)
    for feature in heatmap_data.index:
        for target in heatmap_data.columns:
            value = heatmap_data.loc[feature, target]
            row = results[
                (results["Feature_display"] == feature)
                & (results["Target"] == target)
            ].iloc[0]
            if row["Interpretation"].startswith("Definitional"):
                label = f"{value:.2f}*"
            elif row["Significant_FDR_0.05"]:
                label = f"{value:.2f}**"
            else:
                label = f"{value:.2f}"
            annotation.loc[feature, target] = label

    plt.figure(figsize=(7.2, 8.4))
    sns.heatmap(
        heatmap_data,
        annot=annotation,
        fmt="",
        cmap="vlag",
        center=0,
        vmin=-0.6,
        vmax=0.6,
        linewidths=0.6,
        linecolor="white",
        cbar_kws={"label": "Spearman rho"},
    )
    plt.xlabel("")
    plt.ylabel("")
    plt.title("Final radiomics features versus age and ROI area", pad=12)
    plt.figtext(
        0.01,
        0.01,
        "* Definitional identity; ** FDR-adjusted P < 0.05",
        ha="left",
        va="bottom",
        fontsize=8,
        color="#44515A",
    )
    plt.tight_layout(rect=(0, 0.035, 1, 0.98))
    plt.savefig(OUTPUT_PNG, dpi=300, bbox_inches="tight")
    plt.close()

    ONLINE_XLSX.write_bytes(OUTPUT_XLSX.read_bytes())
    ONLINE_PNG.write_bytes(OUTPUT_PNG.read_bytes())

    print(f"Saved analysis workbook: {OUTPUT_XLSX}")
    print(f"Saved heatmap: {OUTPUT_PNG}")
    print(f"Copied workbook to Supplementary: {ONLINE_XLSX}")
    print(f"Copied heatmap to Supplementary: {ONLINE_PNG}")
    print("\nCorrelation summary:")
    print(
        results[
            [
                "Feature",
                "Target",
                "N",
                "Spearman_rho",
                "CI95_low",
                "CI95_high",
                "P_raw",
                "P_FDR",
                "Interpretation",
            ]
        ].to_string(index=False)
    )


if __name__ == "__main__":
    main()
