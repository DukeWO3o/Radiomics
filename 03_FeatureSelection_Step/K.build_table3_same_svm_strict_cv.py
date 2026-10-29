from __future__ import annotations

import argparse
import importlib.util
import json
import os
import sys
import tempfile
from pathlib import Path

import joblib
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sklearn.metrics import roc_auc_score
from statsmodels.stats.multitest import multipletests


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))
from pipeline_paths import (
    DEFAULT_MODEL_DATA_931,
    OUTPUT_ROOT,
    SUPPLEMENTARY_OUTPUT_ROOT,
)

DEFAULT_STRICT_CV_ROOT = OUTPUT_ROOT
DEFAULT_DATA_FILE = DEFAULT_MODEL_DATA_931
DEFAULT_OUTPUT_XLSX = SUPPLEMENTARY_OUTPUT_ROOT / "Table3_SameSVM_StrictCV.xlsx"
DEFAULT_OUTPUT_JSON = SUPPLEMENTARY_OUTPUT_ROOT / "Table3_SameSVM_StrictCV.json"
DEFAULT_OR6_OUTPUT = (
    SUPPLEMENTARY_OUTPUT_ROOT / "Online_Resource_6_same_SVM_modality_comparison.xlsx"
)
PICKLE_HELPER_SCRIPT = (
    Path(__file__).resolve().parent / "strict_cv_pickle_helpers.py"
)

MODALITIES = {
    "Combined": {
        "directory": "Combined",
        "model_label": "Combined (SVM)",
        "predictors": "4 clinical + 8 radiomics",
    },
    "Clinical-only": {
        "directory": "Clinical",
        "model_label": "Clinical-only (SVM)",
        "predictors": "4 clinical",
    },
    "Radiomics-only": {
        "directory": "Radiomics",
        "model_label": "Radiomics-only (SVM)",
        "predictors": "8 radiomics",
    },
}

OUTPUT_COLUMNS = [
    "Model",
    "Predictors",
    "AUC",
    "AUC_95CI_Low",
    "AUC_95CI_High",
    "Accuracy",
    "Precision",
    "Recall",
    "Specificity",
    "PPV",
    "NPV",
    "F1",
    "BalancedAcc",
    "MCC",
    "Brier",
    "CalibrationIntercept",
    "CalibrationSlope",
    "NetBenefit_0.3",
    "NetBenefit_0.5",
    "NetBenefit_0.7",
    "P_vs_Combined_FDR",
]

METRIC_COLUMNS = {
    "AUC": "Valid_AUC",
    "AUC_95CI_Low": "Valid_AUC_95CI_Low",
    "AUC_95CI_High": "Valid_AUC_95CI_High",
    "Accuracy": "Valid_Accuracy",
    "Precision": "Valid_Precision",
    "Recall": "Valid_Recall",
    "Specificity": "Valid_Specificity",
    "PPV": "Valid_PPV",
    "NPV": "Valid_NPV",
    "F1": "Valid_F1",
    "BalancedAcc": "Valid_BalancedAcc",
    "MCC": "Valid_MCC",
    "Brier": "Valid_Brier",
    "CalibrationIntercept": "Valid_CalibrationIntercept",
    "CalibrationSlope": "Valid_CalibrationSlope",
    "NetBenefit_0.3": "Valid_NetBenefit_0.3",
    "NetBenefit_0.5": "Valid_NetBenefit_0.5",
    "NetBenefit_0.7": "Valid_NetBenefit_0.7",
}


def load_pipeline_module():
    spec = importlib.util.spec_from_file_location(
        "_strict_cv_pipeline_for_table3",
        PICKLE_HELPER_SCRIPT,
    )
    if spec is None or spec.loader is None:
        raise ImportError(
            f"Cannot load pipeline helper: {PICKLE_HELPER_SCRIPT}"
        )

    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)


    setattr(sys.modules["__main__"], "StrictPreprocessSelector", module.StrictPreprocessSelector)
    setattr(sys.modules["__main__"], "ComBatTransformer", module.ComBatTransformer)
    return module


def resolve_modality_paths(strict_cv_root):
    paths = {}
    for modality, settings in MODALITIES.items():
        directory = strict_cv_root / settings["directory"]
        paths[modality] = {
            "directory": directory,
            "metrics": directory / "04_Model_Performance.xlsx",
            "manifest": directory / "00_Validation_Patient_Manifest.pkl",
            "model": directory / "model_SVM.pkl",
        }
    return paths


def load_svm_metrics(metrics_path):
    frame = pd.read_excel(metrics_path)
    model_column = frame.columns[0]
    matched = frame.loc[frame[model_column].astype(str) == "SVM"]
    if len(matched) != 1:
        raise ValueError(f"Expected exactly one SVM row in {metrics_path}")
    return matched.iloc[0]


def validation_predictions(data, manifest_path, model_path):
    manifest = joblib.load(manifest_path)
    required = {"image", "label"}
    if not required.issubset(manifest.columns):
        raise ValueError(f"Manifest must contain {sorted(required)}: {manifest_path}")

    merge_keys = ["image", "label"]
    if {"batch_model"}.issubset(manifest.columns) and {
        "batch_model"
    }.issubset(data.columns):
        merge_keys.append("batch_model")
    merged = manifest.merge(
        data,
        on=merge_keys,
        how="left",
        validate="one_to_one",
    )
    if len(merged) != len(manifest):
        raise ValueError(f"Validation merge changed row count: {manifest_path}")

    model = joblib.load(model_path)
    probabilities = model.predict_proba(merged)[:, 1]
    return manifest, model, probabilities


def get_svm_hyperparameters(model):
    if "classifier" not in model.named_steps:
        raise ValueError("Saved pipeline does not contain a 'classifier' step")
    classifier = model.named_steps["classifier"]

    def json_number(value):
        number = float(value)
        return int(number) if number.is_integer() else number

    return {
        "C": json_number(classifier.C),
        "gamma": json_number(classifier.gamma),
        "kernel": str(classifier.kernel),
    }


def build_rows(metrics_by_modality, fdr_by_comparison):
    rows = []
    for modality, settings in MODALITIES.items():
        metrics = metrics_by_modality[modality]
        row = {
            "Model": settings["model_label"],
            "Predictors": settings["predictors"],
        }
        for output_column, source_column in METRIC_COLUMNS.items():
            row[output_column] = float(metrics[source_column])

        if modality == "Clinical-only":
            row["P_vs_Combined_FDR"] = fdr_by_comparison[
                "Combined vs Clinical-only"
            ]
        elif modality == "Radiomics-only":
            row["P_vs_Combined_FDR"] = fdr_by_comparison[
                "Combined vs Radiomics-only"
            ]
        else:
            row["P_vs_Combined_FDR"] = None
        rows.append(row)
    return pd.DataFrame(rows, columns=OUTPUT_COLUMNS)


def style_worksheet(worksheet):
    thin_gray = Side(style="thin", color="D9D9D9")
    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )
    for row in worksheet.iter_rows():
        for cell in row:
            cell.border = border
            if cell.row == 1:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="top")


def calc_net_benefit(y_true, y_prob, threshold):
    y_true = np.asarray(y_true)
    y_pred = (np.asarray(y_prob) >= threshold).astype(int)
    true_positive = np.sum((y_true == 1) & (y_pred == 1))
    false_positive = np.sum((y_true == 0) & (y_pred == 1))
    n = len(y_true)
    return (
        true_positive / n
        - false_positive / n * threshold / (1 - threshold)
    )


def build_dca_curve_data(labels, predictions):
    thresholds = np.linspace(0.0, 1.0, 101)
    prevalence = float(np.mean(labels))
    curve_data = pd.DataFrame({"Threshold": thresholds})
    for modality, probabilities in predictions.items():
        values = []
        for threshold in thresholds:
            if threshold >= 1.0:
                values.append(0.0)
            else:
                values.append(
                    calc_net_benefit(labels, probabilities, threshold)
                )
        curve_data[modality] = values
    curve_data["Treat_all"] = [
        0.0
        if threshold >= 1.0
        else prevalence
        - (1.0 - prevalence) * threshold / (1.0 - threshold)
        for threshold in thresholds
    ]
    curve_data["Treat_none"] = 0.0
    return curve_data


def render_or6_dca(curve_data, output_svg, output_png):
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "DejaVu Sans", "sans-serif"],
            "font.size": 10,
            "svg.fonttype": "none",
        }
    )
    figure, axes = plt.subplots(1, 2, figsize=(12, 5.3))
    panels = [
        ("Clinical-only", "a", "Clinical-only SVM"),
        ("Radiomics-only", "b", "Radiomics-only SVM"),
    ]
    for axis, (modality, panel_label, title) in zip(axes, panels):
        axis.plot(
            curve_data["Threshold"],
            curve_data[modality],
            color="#C62828",
            linewidth=2.2,
            label=title,
        )
        axis.plot(
            curve_data["Threshold"],
            curve_data["Treat_all"],
            color="#777777",
            linestyle="--",
            linewidth=1.5,
            label="Treat all",
        )
        axis.plot(
            curve_data["Threshold"],
            curve_data["Treat_none"],
            color="#202020",
            linewidth=1.3,
            label="Treat none",
        )
        axis.axhline(0.0, color="#202020", linewidth=0.8)
        axis.set_xlim(0.0, 1.0)
        axis.set_ylim(-0.05, 0.65)
        axis.set_xlabel("Threshold probability", fontsize=11)
        axis.set_ylabel("Net benefit", fontsize=11)
        axis.set_title(
            f"{panel_label}. {title}",
            loc="left",
            fontsize=12,
        )
        axis.grid(True, color="#D9D9D9", linewidth=0.7, alpha=0.65)
        axis.legend(frameon=False, loc="upper right", fontsize=9)
        axis.tick_params(labelsize=9)
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)
    figure.tight_layout()
    figure.savefig(
        output_svg,
        dpi=300,
        bbox_inches="tight",
        facecolor="white",
    )
    figure.savefig(
        output_png,
        dpi=300,
        bbox_inches="tight",
        facecolor="white",
    )
    plt.close(figure)


def build_or6_workbook(table, curve_data, output_path, output_png):
    workbook = Workbook()
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.properties.title = (
        "Online Resource 6: same-SVM feature-set comparison"
    )
    numeric_sheet = workbook.active
    numeric_sheet.title = "Model_comparison"
    numeric_sheet.append(OUTPUT_COLUMNS)
    for record in table.itertuples(index=False, name=None):
        numeric_sheet.append(list(record))

    dark_fill = PatternFill("solid", fgColor="1F4E78")
    pale_fill = PatternFill("solid", fgColor="F7F9FA")
    thin_gray = Side(style="thin", color="D9D9D9")
    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )
    for row in numeric_sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            if cell.row == 1:
                cell.fill = dark_fill
                cell.font = Font(bold=True, color="FFFFFF")
            elif cell.row % 2 == 1:
                cell.fill = pale_fill
    numeric_sheet.freeze_panes = "A2"
    numeric_sheet.row_dimensions[1].height = 34
    widths = {
        "A": 24,
        "B": 28,
        "C": 10,
        "D": 14,
        "E": 14,
        "F": 11,
        "G": 12,
        "H": 10,
        "I": 12,
        "J": 10,
        "K": 10,
        "L": 10,
        "M": 14,
        "N": 10,
        "O": 10,
        "P": 20,
        "Q": 18,
        "R": 18,
        "S": 18,
        "T": 18,
        "U": 20,
    }
    for column, width in widths.items():
        numeric_sheet.column_dimensions[column].width = width
    for row in range(2, 5):
        for column in ["C", "D", "E"]:
            numeric_sheet[f"{column}{row}"].number_format = "0.000"

    chart_sheet = workbook.create_sheet("Feature_set_DCA")
    chart_sheet.sheet_view.showGridLines = False
    chart_sheet["A1"] = (
        "Decision curve analysis using the same SVM for all three feature sets"
    )
    chart_sheet["A1"].font = Font(bold=True, size=14, color="1F4E78")
    chart_sheet.merge_cells("A1:N1")
    chart_sheet["A2"] = (
        "Panel a shows the clinical-only SVM and panel b shows the "
        "radiomics-only SVM. The combined SVM curve is shown in main "
        "manuscript Fig. 6B. All curves use the same locked 187-case "
        "internal validation set."
    )
    chart_sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    chart_sheet.merge_cells("A2:N2")
    chart_sheet.row_dimensions[2].height = 34
    image = ExcelImage(output_png)
    image.width = 1050
    image.height = 464
    chart_sheet.add_image(image, "A4")
    chart_sheet["A32"] = (
        "Source: locked validation predictions from the strict-CV "
        "clinical-only, radiomics-only, and combined SVM pipelines. "
        "The displayed thresholds range from 0 to 1 in increments of 0.01."
    )
    chart_sheet["A32"].alignment = Alignment(wrap_text=True, vertical="top")
    chart_sheet.merge_cells("A32:N32")
    chart_sheet.row_dimensions[32].height = 32

    data_sheet = workbook.create_sheet("DCA_curve_data")
    data_sheet.append(list(curve_data.columns))
    for record in curve_data.itertuples(index=False, name=None):
        data_sheet.append(list(record))
    for cell in data_sheet[1]:
        cell.fill = dark_fill
        cell.font = Font(bold=True, color="FFFFFF")
    data_sheet.freeze_panes = "A2"
    for row in data_sheet.iter_rows(min_row=2):
        row[0].number_format = "0.00"
        for cell in row[1:]:
            cell.number_format = "0.000"

    for worksheet in workbook.worksheets:
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
    workbook.save(output_path)


def atomic_save_workbook(workbook, output_path):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}_",
        suffix=".xlsx",
        dir=output_path.parent,
    )
    os.close(handle)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def atomic_save_json(payload, output_path):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}_",
        suffix=".json",
        dir=output_path.parent,
        text=True,
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(handle, "w", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, indent=2)
            stream.write("\n")
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--strict-cv-root",
        type=Path,
        default=DEFAULT_STRICT_CV_ROOT,
        help="Directory containing the three strict-CV modality outputs.",
    )
    parser.add_argument(
        "--data-file",
        type=Path,
        default=DEFAULT_DATA_FILE,
        help="Merged clinical and radiomics workbook.",
    )
    parser.add_argument(
        "--output-xlsx",
        type=Path,
        default=DEFAULT_OUTPUT_XLSX,
        help="Destination XLSX path.",
    )
    parser.add_argument(
        "--output-json",
        type=Path,
        default=DEFAULT_OUTPUT_JSON,
        help="Destination JSON path.",
    )
    parser.add_argument(
        "--or6-output",
        type=Path,
        default=DEFAULT_OR6_OUTPUT,
        help="Destination Online Resource 6 workbook.",
    )
    parser.add_argument(
        "--write-or6",
        action="store_true",
        help="Also rebuild Online Resource 6 from the same locked SVM outputs.",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    strict_cv_root = args.strict_cv_root.resolve()
    data_file = args.data_file.resolve()

    pipeline_module = load_pipeline_module()
    modality_paths = resolve_modality_paths(strict_cv_root)
    data = pd.read_excel(data_file).reset_index(drop=True)

    metrics_by_modality = {}
    predictions = {}
    manifests = {}
    models = {}

    for modality, paths in modality_paths.items():
        metrics = load_svm_metrics(paths["metrics"])
        manifest, model, probabilities = validation_predictions(
            data,
            paths["manifest"],
            paths["model"],
        )
        observed_auc = float(roc_auc_score(manifest["label"], probabilities))
        recorded_auc = float(metrics["Valid_AUC"])
        if not np.isclose(observed_auc, recorded_auc, rtol=0.0, atol=1e-10):
            raise ValueError(
                f"{modality} validation AUC mismatch: "
                f"recomputed={observed_auc:.12f}, recorded={recorded_auc:.12f}"
            )

        metrics_by_modality[modality] = metrics
        predictions[modality] = probabilities
        manifests[modality] = manifest
        models[modality] = model

    reference_manifest = manifests["Combined"].reset_index(drop=True)
    for modality in MODALITIES:
        current_manifest = manifests[modality].reset_index(drop=True)
        if not reference_manifest.equals(current_manifest):
            raise ValueError(
                f"{modality} does not use the same locked validation cohort/order"
            )

    labels = reference_manifest["label"].to_numpy()
    comparisons = [
        ("Combined vs Clinical-only", "Combined", "Clinical-only"),
        ("Combined vs Radiomics-only", "Combined", "Radiomics-only"),
        ("Clinical-only vs Radiomics-only", "Clinical-only", "Radiomics-only"),
    ]
    raw_p_values = {
        comparison: float(
            pipeline_module.delong_pvalue(
                labels,
                predictions[first],
                predictions[second],
            )
        )
        for comparison, first, second in comparisons
    }
    _, adjusted, _, _ = multipletests(
        list(raw_p_values.values()),
        method="fdr_bh",
    )
    adjusted_p_values = {
        comparison: float(value)
        for comparison, value in zip(raw_p_values, adjusted)
    }

    table = build_rows(metrics_by_modality, adjusted_p_values)
    curve_data = build_dca_curve_data(labels, predictions)
    for modality, threshold in [
        ("Combined", "0.3"),
        ("Clinical-only", "0.3"),
        ("Radiomics-only", "0.3"),
        ("Combined", "0.5"),
        ("Clinical-only", "0.5"),
        ("Radiomics-only", "0.5"),
        ("Combined", "0.7"),
        ("Clinical-only", "0.7"),
        ("Radiomics-only", "0.7"),
    ]:
        threshold_value = float(threshold)
        curve_index = int(round(threshold_value * 100))
        recomputed = float(curve_data.loc[curve_index, modality])
        recorded = float(
            metrics_by_modality[modality][
                f"Valid_NetBenefit_{threshold}"
            ]
        )
        if not np.isclose(recomputed, recorded, rtol=0.0, atol=1e-12):
            raise ValueError(
                f"{modality} net benefit mismatch at {threshold}: "
                f"recomputed={recomputed:.12f}, recorded={recorded:.12f}"
            )

    workbook = Workbook()
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.properties.title = "Table 3: same-SVM feature-set comparison"
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append(OUTPUT_COLUMNS)
    for record in table.itertuples(index=False, name=None):
        worksheet.append(list(record))
    style_worksheet(worksheet)

    combined_model = models["Combined"]
    json_payload = {
        "classifier": "SVM",
        "hyperparameters": get_svm_hyperparameters(combined_model),
        "validation_n": int(len(reference_manifest)),
        "pairwise_raw_p": raw_p_values,
        "pairwise_fdr_p": adjusted_p_values,
    }

    output_xlsx = args.output_xlsx.resolve()
    output_json = args.output_json.resolve()
    atomic_save_workbook(workbook, output_xlsx)
    atomic_save_json(json_payload, output_json)

    if args.write_or6:
        or6_output = args.or6_output.resolve()
        or6_svg = strict_cv_root / "OR6_DCA_same_SVM.svg"
        or6_png = strict_cv_root / "OR6_DCA_same_SVM.png"
        render_or6_dca(curve_data, or6_svg, or6_png)
        build_or6_workbook(table, curve_data, or6_output, or6_png)
        print(f"Online Resource 6 workbook: {or6_output}")
        print(f"Online Resource 6 DCA SVG: {or6_svg}")
        print(f"Online Resource 6 DCA PNG: {or6_png}")


    saved_workbook = load_workbook(output_xlsx, read_only=True, data_only=True)
    saved_sheet = saved_workbook["Sheet1"]
    if saved_sheet.max_row != 4 or saved_sheet.max_column != len(OUTPUT_COLUMNS):
        raise ValueError("Saved Table 3 has an unexpected range")
    saved_workbook.close()

    print(f"Table 3 workbook: {output_xlsx}")
    print(f"Table 3 metadata: {output_json}")
    print(
        "Validation AUCs: "
        + ", ".join(
            f"{modality}={float(metrics_by_modality[modality]['Valid_AUC']):.6f}"
            for modality in MODALITIES
        )
    )
    print(
        "FDR-adjusted pairwise P values: "
        + ", ".join(
            f"{comparison}={value:.6g}"
            for comparison, value in adjusted_p_values.items()
        )
    )


if __name__ == "__main__":
    main()
