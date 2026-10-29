from __future__ import annotations

import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))
from pipeline_paths import (
    APP_MODEL_ROOT,
    COMBINED_OUTPUT_ROOT,
    OUTPUT_ROOT,
    REPEATED_CV_OUTPUT_ROOT,
    SUPPLEMENTARY_OUTPUT_ROOT,
)

STRICT_ROOT = OUTPUT_ROOT
COMBINED_DIR = COMBINED_OUTPUT_ROOT
SUPPLEMENTARY_DIR = SUPPLEMENTARY_OUTPUT_ROOT
OR4_PATH = SUPPLEMENTARY_DIR / "Online_Resource_4_model_reproducibility.xlsx"
OR5_PATH = SUPPLEMENTARY_DIR / "Online_Resource_5_strictCV_DeLong_p_values.xlsx"
OR8_PATH = SUPPLEMENTARY_DIR / "Online_Resource_8_5x5_repeated_CV.xlsx"
REPEATED_CV_PATH = (
    REPEATED_CV_OUTPUT_ROOT / "01_5x5_Repeated_CV_Results.xlsx"
)
APP_MODELS_DIR = APP_MODEL_ROOT


def read_table(path: Path, default_first_header: str | None = None):
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    headers = list(rows[0])
    if headers[0] is None and default_first_header:
        headers[0] = default_first_header
    records = {}
    for row in rows[1:]:
        if row[0] is None:
            continue
        records[str(row[0])] = {
            header: value
            for header, value in zip(headers, row)
            if header is not None
        }
    workbook.close()
    return headers, records


def make_backup_dir() -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = (
        OUTPUT_ROOT / "Backups" / f"StrictCV_rerun_{stamp}"
    )
    backup_dir.mkdir(parents=True, exist_ok=False)
    return backup_dir


def backup_file(path: Path, backup_dir: Path) -> None:
    if path.exists():
        shutil.copy2(path, backup_dir / path.name)


def replace_sheet_rows(worksheet, records) -> None:
    headers = [
        str(cell.value) if cell.value is not None else None
        for cell in worksheet[1]
    ]
    model_column = headers.index("Model") + 1
    for row_index in range(2, worksheet.max_row + 1):
        model = worksheet.cell(row_index, model_column).value
        if model is None or str(model) not in records:
            continue
        record = records[str(model)]
        for column_index, header in enumerate(headers, start=1):
            if header is None or header == "Model":
                continue
            worksheet.cell(
                row_index,
                column_index,
                record.get(header) if header in record else None,
            )


def update_or4() -> None:
    performance_headers, performance = read_table(
        COMBINED_DIR / "04_Model_Performance.xlsx",
        default_first_header="Model",
    )
    _, parameters = read_table(COMBINED_DIR / "04c_Best_Hyperparameters.xlsx")
    _, frequencies = read_table(COMBINED_DIR / "03_Selection_Frequency.xlsx")
    model_config = json.loads(
        (COMBINED_DIR / "model_config.json").read_text(encoding="utf-8")
    )
    final_features = set(model_config["final_features"])

    workbook = load_workbook(OR4_PATH)
    replace_sheet_rows(workbook["Model_Performance"], performance)
    replace_sheet_rows(workbook["Best_Hyperparameters"], parameters)

    reproducibility = workbook["Final_Feature_Reproducibility"]
    headers = {
        str(cell.value): cell.column
        for cell in reproducibility[1]
        if cell.value is not None
    }
    feature_column = headers["Feature"]
    for row_index in range(2, reproducibility.max_row + 1):
        feature = reproducibility.cell(row_index, feature_column).value
        if feature is None or str(feature) not in frequencies:
            continue
        record = frequencies[str(feature)]
        reproducibility.cell(
            row_index,
            headers["Selected_in_folds"],
            record.get("Selected_in_folds"),
        )
        reproducibility.cell(
            row_index,
            headers["Selection_frequency"],
            record.get("Frequency"),
        )
        reproducibility.cell(
            row_index,
            headers["Protected"],
            record.get("Protected"),
        )
        reproducibility.cell(
            row_index,
            headers["In_final_model"],
            str(feature) in final_features,
        )

    workbook.properties.lastModifiedBy = ""
    workbook.save(OR4_PATH)


def update_or5() -> None:
    shutil.copy2(
        COMBINED_DIR / "11_DeLong_Pvalues_FDRcorrected.xlsx",
        OR5_PATH,
    )


def update_or8() -> None:
    shutil.copy2(REPEATED_CV_PATH, OR8_PATH)


def update_app_models(backup_dir: Path) -> None:
    APP_MODELS_DIR.mkdir(parents=True, exist_ok=True)
    model_backup = backup_dir / "app_models"
    model_backup.mkdir()
    for path in sorted(APP_MODELS_DIR.iterdir()):
        if path.is_file() and (
            path.name == "model_config.json" or path.name.startswith("model_")
        ):
            shutil.copy2(path, model_backup / path.name)

    for path in sorted(COMBINED_DIR.glob("model_*.pkl")):
        shutil.copy2(path, APP_MODELS_DIR / path.name)
    shutil.copy2(
        COMBINED_DIR / "model_config.json",
        APP_MODELS_DIR / "model_config.json",
    )


def verify() -> None:
    stra = json.loads(
        (COMBINED_DIR / "model_config.json").read_text(encoding="utf-8")
    )
    app = json.loads(
        (APP_MODELS_DIR / "model_config.json").read_text(encoding="utf-8")
    )
    if stra["final_features"] != app["final_features"]:
        raise ValueError("Application and strict-CV final feature lists differ")

    for source in sorted(COMBINED_DIR.glob("model_*.pkl")):
        target = APP_MODELS_DIR / source.name
        if source.read_bytes() != target.read_bytes():
            raise ValueError(f"Application model copy differs: {source.name}")

    workbook = load_workbook(OR4_PATH, read_only=True, data_only=True)
    headers = [
        cell.value for cell in next(workbook["Model_Performance"].iter_rows())
    ]
    if len(headers) != 38:
        raise ValueError("Online Resource 4 model-performance columns changed")
    workbook.close()

    workbook = load_workbook(OR5_PATH, read_only=True, data_only=True)
    if workbook.active.max_row != 10 or workbook.active.max_column != 10:
        raise ValueError("Online Resource 5 DeLong matrix has an unexpected shape")
    workbook.close()

    workbook = load_workbook(OR8_PATH, read_only=True, data_only=True)
    if "Summary" not in workbook.sheetnames or "Final_Parameters" not in workbook.sheetnames:
        raise ValueError("Online Resource 8 is missing expected sheets")
    workbook.close()


def main() -> None:
    backup_dir = make_backup_dir()
    backup_file(OR4_PATH, backup_dir)
    backup_file(OR5_PATH, backup_dir)
    backup_file(OR8_PATH, backup_dir)
    update_or4()
    update_or5()
    update_or8()
    update_app_models(backup_dir)
    verify()
    print(f"Backup: {backup_dir}")
    print(f"Updated: {OR4_PATH}")
    print(f"Updated: {OR5_PATH}")
    print(f"Updated: {OR8_PATH}")
    print(f"Updated application models: {APP_MODELS_DIR}")


if __name__ == "__main__":
    main()
