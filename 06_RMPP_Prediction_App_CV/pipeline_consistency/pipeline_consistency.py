from __future__ import annotations

import hashlib
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


BASE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = BASE_DIR.parents[1]
sys.path.insert(0, str(PROJECT_DIR))
from pipeline_paths import APP_DEPLOYMENT_OUTPUT_ROOT, SUPPLEMENTARY_OUTPUT_ROOT

SOURCE_WORKBOOK = (
    APP_DEPLOYMENT_OUTPUT_ROOT / "Research_vs_Deployment_Pipeline_Consistency.xlsx"
)
METRICS_CSV = (
    BASE_DIR.parent / "unet_validation_output" / "01_Test_Case_Metrics.csv"
)
TARGET_WORKBOOK = (
    SUPPLEMENTARY_OUTPUT_ROOT / "Online_Resource_9_U-Net_and_pipeline_validation.xlsx"
)
MODEL_CONFIG = BASE_DIR.parent / "models" / "model_config.json"
U_NET_CHECKPOINT = BASE_DIR.parent / "models" / "cxr_unet_best.pth"
RESEARCH_RADIOMICS_CONFIG = (
    PROJECT_DIR / "01_PreProcess" / "radiomics_config.yaml"
)
DEPLOYMENT_RADIOMICS_CONFIG = (
    BASE_DIR.parent / "config" / "radiomics_config.yaml"
)


def file_digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def append_readme_rows(worksheet) -> None:
    entries = [
        (
            "Pipeline_Consistency",
            "SVM probability agreement and final radiomics-feature agreement.",
        ),
        (
            "Pipeline_Case_Level",
            "Case-level probabilities and final radiomics-feature values.",
        ),
    ]
    redundant_labels = {
        "Research-to-deployment consistency",
        "Pipeline consistency cohort",
        "Pipeline configuration",
    }
    for row in range(worksheet.max_row, 0, -1):
        if str(worksheet.cell(row, 1).value or "") in redundant_labels:
            worksheet.delete_rows(row)
    existing_rows = {
        str(worksheet.cell(row, 1).value): row
        for row in range(1, worksheet.max_row + 1)
        if worksheet.cell(row, 1).value
    }
    run_summary = (
        f"Pipeline consistency rerun {datetime.now().astimezone().date().isoformat()} "
        f"with App model config SHA256 {file_digest(MODEL_CONFIG)} and "
        f"U-Net checkpoint SHA256 {file_digest(U_NET_CHECKPOINT)}. "
        f"Research radiomics config SHA256 "
        f"{file_digest(RESEARCH_RADIOMICS_CONFIG)}; deployment radiomics "
        f"config SHA256 {file_digest(DEPLOYMENT_RADIOMICS_CONFIG)}."
    )
    configuration_row = existing_rows.get("Configuration")
    if configuration_row is None:
        configuration_row = worksheet.max_row + 1
        worksheet.cell(configuration_row, 1, "Configuration")
    worksheet.cell(
        configuration_row,
        2,
        "Research and deployment extraction settings were aligned before "
        "the reported analysis. "
        + run_summary,
    )
    start_row = worksheet.max_row + 2
    for label, description in entries:
        if label in existing_rows:
            worksheet.cell(existing_rows[label], 2, description)
        else:
            worksheet.cell(start_row, 1, label)
            worksheet.cell(start_row, 2, description)
            start_row += 1


def write_probability_table(worksheet, probability_agreement: pd.DataFrame) -> None:
    start_row = worksheet.max_row + 3
    worksheet.cell(start_row, 1, "Probability agreement")
    worksheet.cell(start_row, 1).font = Font(bold=True, size=11)
    headers = [
        "Comparison",
        "Mean difference (deployment - research)",
        "Mean absolute difference",
        "Bland-Altman lower",
        "Bland-Altman upper",
        "ICC(2,1)",
        "Pearson r",
        "Classification agreement",
        "Cohen kappa",
    ]
    columns = [
        "Comparison",
        "Mean_Difference_Y_minus_X",
        "Mean_Absolute_Difference",
        "Bland_Altman_Lower",
        "Bland_Altman_Upper",
        "ICC_2_1",
        "Pearson_r",
        "Classification_Agreement",
        "Cohen_Kappa",
    ]
    header_row = start_row + 1
    for column_index, header in enumerate(headers, start=1):
        worksheet.cell(header_row, column_index, header)
    for row_index, (_, record) in enumerate(probability_agreement.iterrows(), start=2):
        for column_index, column in enumerate(columns, start=1):
            worksheet.cell(header_row + row_index, column_index, record[column])


def write_feature_table(worksheet, feature_agreement: pd.DataFrame) -> None:
    start_row = worksheet.max_row + 3
    worksheet.cell(start_row, 1, "Final radiomics-feature agreement")
    worksheet.cell(start_row, 1).font = Font(bold=True, size=11)
    headers = [
        "Comparison",
        "Feature",
        "Mean difference (deployment - research)",
        "Mean absolute difference",
        "Bland-Altman lower",
        "Bland-Altman upper",
        "ICC(2,1)",
        "Pearson r",
        "Spearman rho",
    ]
    columns = [
        "Comparison",
        "Feature",
        "Mean_Difference_Y_minus_X",
        "Mean_Absolute_Difference",
        "Bland_Altman_Lower",
        "Bland_Altman_Upper",
        "ICC_2_1",
        "Pearson_r",
        "Spearman_rho",
    ]
    header_row = start_row + 1
    for column_index, header in enumerate(headers, start=1):
        worksheet.cell(header_row, column_index, header)
    for row_index, (_, record) in enumerate(feature_agreement.iterrows(), start=2):
        for column_index, column in enumerate(columns, start=1):
            worksheet.cell(header_row + row_index, column_index, record[column])


def write_excluded_cases(worksheet, excluded: pd.DataFrame) -> None:
    start_row = worksheet.max_row + 3
    worksheet.cell(start_row, 1, "Excluded U-Net test cases")
    worksheet.cell(start_row, 1).font = Font(bold=True, size=11)
    headers = ["Case_ID", "Dice", "HD95_mm", "Exclusion reason"]
    columns = ["Case_ID", "Dice", "HD95_mm", "Exclusion_Reason"]
    header_row = start_row + 1
    for column_index, header in enumerate(headers, start=1):
        worksheet.cell(header_row, column_index, header)
    for row_index, (_, record) in enumerate(excluded.iterrows(), start=2):
        for column_index, column in enumerate(columns, start=1):
            worksheet.cell(header_row + row_index, column_index, record[column])


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_name in ("Pipeline_Consistency", "Pipeline_Case_Level"):
        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.row == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                elif cell.value is not None:
                    cell.font = body_font
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = border
        for column in worksheet.columns:
            maximum = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column
            )
            worksheet.column_dimensions[column[0].column_letter].width = min(
                max(maximum + 2, 12), 52
            )
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False

        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
    workbook.save(path)


def main() -> None:
    summary = pd.read_excel(SOURCE_WORKBOOK, sheet_name="Summary")
    probability = pd.read_excel(SOURCE_WORKBOOK, sheet_name="Probability_Agreement")
    features = pd.read_excel(SOURCE_WORKBOOK, sheet_name="Feature_Agreement")
    cases = pd.read_excel(SOURCE_WORKBOOK, sheet_name="Case_Level_Results")
    excluded = pd.read_excel(SOURCE_WORKBOOK, sheet_name="Excluded_Cases")
    metric_names = pd.read_csv(METRICS_CSV, encoding="utf-8-sig")["Image"].astype(str)
    mapping = {
        name: f"UNET-{index:03d}"
        for index, name in enumerate(sorted(set(metric_names)), start=1)
    }
    if "Image" in cases.columns:
        cases.insert(0, "Case_ID", cases["Image"].map(mapping))
        cases = cases.drop(columns=["Image"])
    if "Image" in excluded.columns:
        excluded.insert(0, "Case_ID", excluded["Image"].map(mapping))
        excluded = excluded.drop(columns=["Image"])
    if cases["Case_ID"].isna().any() or excluded["Case_ID"].isna().any():
        raise ValueError("Some pipeline cases could not be mapped to anonymized IDs.")

    workbook = load_workbook(TARGET_WORKBOOK)
    append_readme_rows(workbook["README"])

    for sheet_name in ("Pipeline_Consistency", "Pipeline_Case_Level"):
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]
    consistency = workbook.create_sheet("Pipeline_Consistency")
    consistency.append(["Metric", "Value"])
    for _, row in summary.iterrows():
        consistency.append([row["Metric"], row["Value"]])
    write_probability_table(consistency, probability)
    write_feature_table(consistency, features)
    write_excluded_cases(consistency, excluded)

    case_sheet = workbook.create_sheet("Pipeline_Case_Level")
    case_sheet.append(cases.columns.tolist())
    for row in cases.itertuples(index=False, name=None):
        case_sheet.append(list(row))

    workbook.save(TARGET_WORKBOOK)
    style_workbook(TARGET_WORKBOOK)
    check = load_workbook(TARGET_WORKBOOK, read_only=True, data_only=True)
    strings = [
        str(value)
        for worksheet in check.worksheets
        for row in worksheet.iter_rows(values_only=True)
        for value in row
        if value is not None
    ]
    combined = "\n".join(strings)
    leaks = [
        name
        for name in mapping
        if name in combined or Path(name).stem in combined
    ]
    if leaks:
        raise ValueError(f"Privacy audit failed: {leaks[:5]}")
    print(f"Updated: {TARGET_WORKBOOK}")
    print(f"Sheets: {load_workbook(TARGET_WORKBOOK, read_only=True).sheetnames}")


if __name__ == "__main__":
    main()
