from __future__ import annotations

import sys
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


BASE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = BASE_DIR.parents[1]
sys.path.insert(0, str(PROJECT_DIR))
from pipeline_paths import SUPPLEMENTARY_OUTPUT_ROOT
SOURCE_WORKBOOK = (
    BASE_DIR.parent / "benchmark_output" / "End_to_End_Timing_Results.xlsx"
)
TARGET_WORKBOOK = (
    SUPPLEMENTARY_OUTPUT_ROOT / "Online_Resource_9_U-Net_and_pipeline_validation.xlsx"
)
TEST_METRICS = (
    BASE_DIR.parent / "unet_validation_output" / "01_Test_Case_Metrics.csv"
)


def add_readme_rows(worksheet) -> None:
    entries = [
        (
            "End_to_End_Timing",
            "DICOM-to-probability benchmark endpoint, technical success, total "
            "latency, stage timing, and asset-load time.",
        ),
        (
            "End_to_End_Timing_Cases",
            "Per-case median timing over three benchmark repetitions using "
            "BENCH-### identifiers.",
        ),
    ]
    existing = {
        str(worksheet.cell(row, 1).value): row
        for row in range(1, worksheet.max_row + 1)
        if worksheet.cell(row, 1).value
    }
    next_row = worksheet.max_row + 1
    for label, description in entries:
        if label in existing:
            worksheet.cell(existing[label], 2, description)
        else:
            worksheet.cell(next_row, 1, label)
            worksheet.cell(next_row, 2, description)
            next_row += 1


def copy_dataframe(worksheet, dataframe: pd.DataFrame, start_row: int = 1) -> int:
    for column_index, column in enumerate(dataframe.columns, start=1):
        worksheet.cell(start_row, column_index, column)
    for row_index, values in enumerate(
        dataframe.itertuples(index=False, name=None),
        start=start_row + 1,
    ):
        for column_index, value in enumerate(values, start=1):
            worksheet.cell(row_index, column_index, value)
    return start_row + len(dataframe) + 2


def build_timing_sheet(workbook, summary, stages, assets) -> None:
    worksheet = workbook.create_sheet("End_to_End_Timing")
    worksheet.append(["Metric", "Value"])

    protocol_rows = [
        (
            "Primary endpoint",
            "From DICOM bytes available in application memory to finite final "
            "SVM probability in [0,1].",
        ),
        (
            "Excluded",
            "Browser upload, disk read, application startup, Streamlit rendering, "
            "and user idle time.",
        ),
        (
            "Inference hardware",
            "CPU, matching the deployed Streamlit app path.",
        ),
        (
            "Repetition design",
            "93 cases, 3 repetitions each, 2 warm-up runs, randomized execution order.",
        ),
    ]
    for row in protocol_rows:
        worksheet.append(row)

    worksheet.append([])
    worksheet.append(["Benchmark summary", "Value"])
    for row in summary.itertuples(index=False, name=None):
        worksheet.append(list(row))

    worksheet.append([])
    worksheet.append(["Asset load", "Value"])
    for row in assets.itertuples(index=False, name=None):
        worksheet.append(list(row))

    worksheet.append([])
    stage_start = worksheet.max_row + 1
    copy_dataframe(worksheet, stages, start_row=stage_start)


def build_case_sheet(workbook, cases: pd.DataFrame) -> None:
    worksheet = workbook.create_sheet("End_to_End_Timing_Cases")
    copy_dataframe(worksheet, cases)


def style_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    section_font = Font(name="Arial", size=11, bold=True, color="1F1F1F")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    section_labels = {
        "Benchmark summary",
        "Asset load",
        "Stage",
    }
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.font = body_font
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
            if isinstance(cell.value, str) and cell.value in section_labels:
                cell.font = section_font
            if isinstance(cell.value, float):
                cell.number_format = "0.000"
    for column in worksheet.columns:
        maximum = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column
        )
        worksheet.column_dimensions[column[0].column_letter].width = min(
            max(maximum + 2, 12),
            54,
        )
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False


def privacy_audit(workbook) -> None:
    source_names = pd.read_csv(TEST_METRICS, encoding="utf-8-sig")[
        "Image"
    ].astype(str)
    forbidden = set(source_names) | {Path(name).stem for name in source_names}
    strings = [
        str(value)
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows(values_only=True)
        for value in row
        if value is not None
    ]
    combined = "\n".join(strings)
    leaks = [token for token in forbidden if token in combined]
    path_leaks = [
        token
        for token in ("Image_Path", "Mask_Path")
        if token in combined
    ]
    if re.search(r"[A-Za-z]:[\\/]", combined):
        path_leaks.append("absolute_path")
    if leaks or path_leaks:
        raise ValueError(
            f"Privacy audit failed. Name leaks={leaks[:5]}, path leaks={path_leaks}"
        )


def main() -> None:
    summary = pd.read_excel(SOURCE_WORKBOOK, sheet_name="Summary")
    stages = pd.read_excel(SOURCE_WORKBOOK, sheet_name="Stage_Summary")
    assets = pd.read_excel(SOURCE_WORKBOOK, sheet_name="Asset_Load")
    cases = pd.read_excel(SOURCE_WORKBOOK, sheet_name="Case_Summary")
    cases["Case_ID"] = cases["Case_ID"].astype(str).str.replace(
        "UNET-", "BENCH-", regex=False
    )

    workbook = load_workbook(TARGET_WORKBOOK)
    for sheet_name in ("End_to_End_Timing", "End_to_End_Timing_Cases"):
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]

    add_readme_rows(workbook["README"])
    build_timing_sheet(workbook, summary, stages, assets)
    build_case_sheet(workbook, cases)
    style_sheet(workbook["README"])
    style_sheet(workbook["End_to_End_Timing"])
    style_sheet(workbook["End_to_End_Timing_Cases"])
    privacy_audit(workbook)
    workbook.save(TARGET_WORKBOOK)
    print(f"Updated: {TARGET_WORKBOOK}")


if __name__ == "__main__":
    main()
