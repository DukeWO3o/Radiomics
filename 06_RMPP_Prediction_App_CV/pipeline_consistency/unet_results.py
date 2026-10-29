from __future__ import annotations

import sys
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.pagebreak import Break
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


BASE_DIR = Path(__file__).resolve().parent
APP_DIR = BASE_DIR.parent
PROJECT_DIR = APP_DIR.parent
sys.path.insert(0, str(PROJECT_DIR))
from pipeline_paths import APP_DEPLOYMENT_OUTPUT_ROOT, SUPPLEMENTARY_OUTPUT_ROOT

OR9_PATH = (
    SUPPLEMENTARY_OUTPUT_ROOT / "Online_Resource_9_U-Net_and_pipeline_validation.xlsx"
)
METRICS_CSV = APP_DIR / "unet_validation_output" / "01_Test_Case_Metrics.csv"
HISTORY_CSV = APP_DIR / "unet_validation_output" / "03_Training_History.csv"
SPLIT_CSV = APP_DIR / "unet_validation_output" / "02_Split_Manifest.csv"
EXAMPLE_PATH = (
    APP_DEPLOYMENT_OUTPUT_ROOT / "U-Net_Test_Examples_Anonymized.png"
)

sys.path.insert(0, str(APP_DIR))
from train_unet_independent_test import save_example_figure


def anonymized_mapping(images) -> dict[str, str]:
    names = sorted({Path(str(name)).name for name in images})
    return {name: f"UNET-{index:03d}" for index, name in enumerate(names, start=1)}


def style_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in worksheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = header_font
            else:
                cell.font = body_font
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
    for column in worksheet.columns:
        width = max(
            len(str(cell.value)) if cell.value is not None else 0 for cell in column
        )
        worksheet.column_dimensions[column[0].column_letter].width = min(
            max(width + 2, 12), 52
        )
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False


def update_readme(worksheet) -> None:
    updates = {
        "Case_ID": (
            "UNET-001 to UNET-140 provide consistent mapping for the U-Net and "
            "pipeline sheets. The earlier timing benchmark uses BENCH-001 to "
            "BENCH-093."
        ),
        "Cohort": (
            "U-Net validation used 140 independent test cases from the final "
            "931-case cohort; all 140 had paired research reference features."
        ),
    }
    for row in range(1, worksheet.max_row + 1):
        label = str(worksheet.cell(row, 1).value or "")
        if label in updates:
            worksheet.cell(row, 2, updates[label])


def update_validation_sheet(worksheet, metrics, history, split_manifest) -> None:
    worksheet._images.clear()
    worksheet.row_breaks.brk = []
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)

    rows = [
        ("Matched image-mask pairs", len(split_manifest)),
        ("Training images", int((split_manifest["Split"] == "Train").sum())),
        ("Validation images", int((split_manifest["Split"] == "Validation").sum())),
        ("Independent test images", len(metrics)),
        ("One image per patient", "Yes"),
        ("Split level", "Case-level, one radiograph per patient"),
        ("Training epochs completed", int(len(history))),
        ("Best validation epoch", int(history.loc[history["Validation_Dice"].idxmax(), "Epoch"])),
        ("Best validation Dice", float(history["Validation_Dice"].max())),
        (
            "Mean Dice (SD)",
            f"{metrics['Dice'].mean():.3f} ({metrics['Dice'].std(ddof=1):.3f})",
        ),
        (
            "Median Dice (range)",
            f"{metrics['Dice'].median():.3f} "
            f"({metrics['Dice'].min():.3f}-{metrics['Dice'].max():.3f})",
        ),
        ("Mean HD95, mm", float(metrics["HD95_mm"].mean())),
        ("Median HD95, mm", float(metrics["HD95_mm"].median())),
        ("Mean ASSD, mm", float(metrics["ASSD_mm"].mean())),
        ("Median ASSD, mm", float(metrics["ASSD_mm"].median())),
        ("Dice <0.80, n", int((metrics["Dice"] < 0.80).sum())),
        ("Empty prediction, n", int((metrics["Prediction_Area_Pixels"] == 0).sum())),
        (
            "Area-ratio failure, n",
            int(metrics["Failure_Reasons"].fillna("").str.contains("area_ratio").sum()),
        ),
        ("Any segmentation failure, n", int(metrics["Failure"].sum())),
        ("Any segmentation failure, %", float(metrics["Failure"].mean() * 100)),
        ("Mean inference time, ms", float(metrics["Inference_ms"].mean())),
        ("Median inference time, ms", float(metrics["Inference_ms"].median())),
        (
            "Failure criteria",
            "Dice <0.80, empty prediction, or predicted/reference area ratio "
            "<0.5 or >2.0",
        ),
    ]
    for row_index, record in enumerate(rows, start=2):
        worksheet.cell(row_index, 1, record[0])
        worksheet.cell(row_index, 2, record[1])

    save_example_figure(metrics, EXAMPLE_PATH, n_cases=3)
    worksheet.row_breaks.append(Break(id=worksheet.max_row))
    image = XLImage(str(EXAMPLE_PATH))
    image.width = 1660
    image.height = 1116
    worksheet.add_image(image, f"B{worksheet.max_row + 2}")


def relabel_historical_timing(worksheet) -> None:
    for row in range(2, worksheet.max_row + 1):
        value = str(worksheet.cell(row, 1).value or "")
        if value.startswith("UNET-"):
            worksheet.cell(row, 1, value.replace("UNET-", "BENCH-", 1))


def main() -> None:
    metrics = pd.read_csv(METRICS_CSV, encoding="utf-8-sig")
    history = pd.read_csv(HISTORY_CSV, encoding="utf-8-sig")
    split_manifest = pd.read_csv(SPLIT_CSV, encoding="utf-8-sig")
    mapping = anonymized_mapping(metrics["Image"])

    metrics = metrics.copy()
    metrics["Case_ID"] = metrics["Image"].map(mapping)
    if metrics["Case_ID"].isna().any():
        raise ValueError("Some U-Net test cases could not be anonymized.")

    workbook = load_workbook(OR9_PATH)
    update_readme(workbook["README"])
    readme = workbook["README"]
    for row in range(1, readme.max_row + 1):
        if str(readme.cell(row, 1).value or "") == "End_to_End_Timing_Cases":
            readme.cell(
                row,
                2,
                "Per-case median timing over three repetitions for the earlier "
                "93-case technical benchmark using BENCH-### identifiers.",
            )
    update_validation_sheet(
        workbook["U-Net_Validation"],
        metrics,
        history,
        split_manifest,
    )

    old_index = workbook.sheetnames.index("U-Net_Test_Cases")
    del workbook["U-Net_Test_Cases"]
    test_sheet = workbook.create_sheet("U-Net_Test_Cases", old_index)
    columns = [
        "Case_ID",
        "Dice",
        "HD95_mm",
        "ASSD_mm",
        "Reference_Area_Pixels",
        "Prediction_Area_Pixels",
        "Area_Ratio",
        "Failure",
        "Failure_Reasons",
        "Inference_ms",
    ]
    test_sheet.append(columns)
    for row in metrics[columns].itertuples(index=False, name=None):
        test_sheet.append(list(row))
    relabel_historical_timing(workbook["End_to_End_Timing_Cases"])
    style_sheet(workbook["README"])
    style_sheet(test_sheet)
    style_sheet(workbook["U-Net_Validation"])
    workbook.save(OR9_PATH)

    check = load_workbook(OR9_PATH, read_only=True, data_only=True)
    strings = [
        str(value)
        for worksheet in check.worksheets
        for row in worksheet.iter_rows(values_only=True)
        for value in row
        if value is not None
    ]
    combined = "\n".join(strings)
    leaks = [
        name
        for name in mapping
        if name in combined or Path(name).stem in combined
    ]
    path_leaks = [
        token
        for token in ("DCM-RMPP", "01_PNG", "03_Binarized")
        if token in combined
    ]
    if re.search(r"[A-Za-z]:[\\/]", combined):
        path_leaks.append("absolute_path")
    if leaks or path_leaks:
        raise ValueError(
            f"Privacy audit failed. Name leaks={leaks[:5]}, path leaks={path_leaks}"
        )
    print(f"Updated U-Net sheets: {OR9_PATH}")
    print(f"Anonymized example: {EXAMPLE_PATH}")


if __name__ == "__main__":
    main()
