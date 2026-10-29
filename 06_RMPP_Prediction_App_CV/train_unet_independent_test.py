from __future__ import annotations

import json
import math
import random
import shutil
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import torch.optim as optim
from PIL import Image, ImageDraw, ImageFont
from scipy.ndimage import (
    binary_erosion,
    distance_transform_edt,
    generate_binary_structure,
)
from sklearn.model_selection import train_test_split
from torch.utils.data import DataLoader, Dataset
from torchvision import transforms
from tqdm import tqdm

from unet_model import SimpleUNet


torch.backends.cudnn.enabled = False
PIPELINE_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PIPELINE_ROOT))
from pipeline_paths import DEFAULT_MODEL_DATA_931, IMAGE_ROOT


class Config:
    COHORT_FILE = DEFAULT_MODEL_DATA_931
    DATASET_DIRS = {
        1: IMAGE_ROOT / "RMPP",
        0: IMAGE_ROOT / "nonRMPP",
    }
    DATASET_NAMES = {1: "RMPP", 0: "nonRMPP"}
    OUTPUT_DIR = Path(__file__).resolve().parent / "unet_validation_output"
    MODEL_DIR = Path(__file__).resolve().parent / "models"
    BEST_MODEL_PATH = MODEL_DIR / "cxr_unet_best_independent.pth"
    DEPLOY_MODEL_PATHS = (MODEL_DIR / "cxr_unet_best.pth",)

    IMG_SIZE = 512
    BATCH_SIZE = 8
    LEARNING_RATE = 1e-4
    EPOCHS = 50
    EARLY_STOP_PATIENCE = 10
    TRAIN_FRACTION = 0.70
    VAL_FRACTION = 0.15
    TEST_FRACTION = 0.15
    RANDOM_STATE = 42
    DEVICE = "cuda" if torch.cuda.is_available() else "cpu"


    MM_PER_MODEL_PIXEL = 1.0
    FAILURE_DICE_THRESHOLD = 0.80
    FAILURE_AREA_RATIO_LOW = 0.50
    FAILURE_AREA_RATIO_HIGH = 2.00


def seed_everything(seed: int) -> None:
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed)


class CXRDataset(Dataset):
    def __init__(self, records, img_size=512):
        self.records = list(records)
        self.image_transform = transforms.Compose(
            [
                transforms.Resize(
                    (img_size, img_size),
                    interpolation=transforms.InterpolationMode.BILINEAR,
                ),
                transforms.ToTensor(),
            ]
        )
        self.mask_transform = transforms.Compose(
            [
                transforms.Resize(
                    (img_size, img_size),
                    interpolation=transforms.InterpolationMode.NEAREST,
                ),
                transforms.ToTensor(),
            ]
        )

    def __len__(self):
        return len(self.records)

    def __getitem__(self, index):
        record = self.records[index]
        image = Image.open(record["Image_Path"]).convert("L")
        mask = Image.open(record["Mask_Path"]).convert("L")
        image = self.image_transform(image)
        mask = (self.mask_transform(mask) > 0).float()
        return image, mask


class DiceBCELoss(nn.Module):
    def __init__(self):
        super().__init__()
        self.bce = nn.BCEWithLogitsLoss()

    def forward(self, logits, targets, smooth=1.0):
        bce = self.bce(logits, targets)
        probabilities = torch.sigmoid(logits).view(-1)
        targets = targets.view(-1)
        intersection = (probabilities * targets).sum()
        dice_loss = 1.0 - (
            (2.0 * intersection + smooth)
            / (probabilities.sum() + targets.sum() + smooth)
        )
        return bce + dice_loss


def per_image_dice(logits, targets, threshold=0.5, smooth=1e-5):
    predictions = (torch.sigmoid(logits) > threshold).float()
    dims = tuple(range(1, predictions.ndim))
    intersection = (predictions * targets).sum(dim=dims)
    denominator = predictions.sum(dim=dims) + targets.sum(dim=dims)
    return ((2.0 * intersection + smooth) / (denominator + smooth)).detach().cpu()


def surface_distances(reference, prediction):
    structure = generate_binary_structure(2, 1)
    reference = reference.astype(bool)
    prediction = prediction.astype(bool)
    reference_boundary = reference & ~binary_erosion(reference, structure=structure)
    prediction_boundary = prediction & ~binary_erosion(prediction, structure=structure)

    if not reference_boundary.any() or not prediction_boundary.any():
        return np.asarray([np.inf], dtype=float), np.asarray([np.inf], dtype=float)

    distance_to_reference = distance_transform_edt(~reference_boundary)
    distance_to_prediction = distance_transform_edt(~prediction_boundary)
    prediction_to_reference = distance_to_reference[prediction_boundary]
    reference_to_prediction = distance_to_prediction[reference_boundary]
    return prediction_to_reference, reference_to_prediction


def segmentation_metrics(reference, prediction, mm_per_pixel):
    reference = reference.astype(bool)
    prediction = prediction.astype(bool)
    reference_area = int(reference.sum())
    prediction_area = int(prediction.sum())

    if reference_area == 0 and prediction_area == 0:
        dice = 1.0
    else:
        intersection = np.logical_and(reference, prediction).sum()
        dice = 2.0 * intersection / (reference_area + prediction_area)

    pred_to_ref, ref_to_pred = surface_distances(reference, prediction)
    all_distances = np.concatenate([pred_to_ref, ref_to_pred])
    if np.isfinite(all_distances).any():
        hd95 = float(np.percentile(all_distances, 95) * mm_per_pixel)
        assd = float(
            (
                np.mean(pred_to_ref[np.isfinite(pred_to_ref)])
                + np.mean(ref_to_pred[np.isfinite(ref_to_pred)])
            )
            / 2.0
            * mm_per_pixel
        )
    else:
        hd95 = math.nan
        assd = math.nan

    area_ratio = (
        prediction_area / reference_area if reference_area > 0 else math.nan
    )
    failure_reasons = []
    if prediction_area == 0:
        failure_reasons.append("empty_prediction")
    if dice < Config.FAILURE_DICE_THRESHOLD:
        failure_reasons.append("dice_below_0.80")
    if math.isfinite(area_ratio) and (
        area_ratio < Config.FAILURE_AREA_RATIO_LOW
        or area_ratio > Config.FAILURE_AREA_RATIO_HIGH
    ):
        failure_reasons.append("area_ratio_outside_0.5-2.0")

    return {
        "Dice": float(dice),
        "HD95_mm": hd95,
        "ASSD_mm": assd,
        "Reference_Area_Pixels": reference_area,
        "Prediction_Area_Pixels": prediction_area,
        "Area_Ratio": area_ratio,
        "Failure": bool(failure_reasons),
        "Failure_Reasons": "; ".join(failure_reasons),
    }


def validate(model, loader, criterion):
    model.eval()
    total_loss = 0.0
    dice_values = []
    with torch.no_grad():
        for images, masks in loader:
            images = images.to(Config.DEVICE)
            masks = masks.to(Config.DEVICE)
            with torch.cuda.amp.autocast(enabled=Config.DEVICE == "cuda"):
                logits = model(images)
                loss = criterion(logits, masks)
            total_loss += loss.item()
            dice_values.extend(per_image_dice(logits, masks).numpy().tolist())
    return total_loss / max(len(loader), 1), float(np.mean(dice_values))


def _index_dataset(root):
    image_files = {
        path.name.casefold(): path
        for path in (root / "01_PNG").iterdir()
        if path.is_file()
    }
    mask_files = {
        path.name.casefold(): path
        for path in (root / "03_Binarized").iterdir()
        if path.is_file()
    }
    return image_files, mask_files


def load_cohort_records():
    cohort = pd.read_excel(Config.COHORT_FILE, sheet_name=0, usecols=["label", "image"])
    indexes = {
        label: _index_dataset(root) for label, root in Config.DATASET_DIRS.items()
    }
    records = []
    fallback_count = 0

    for row in cohort.itertuples(index=False):
        label = int(row.label)
        image_key = Path(str(row.image).strip()).name.casefold()
        candidate_labels = [label, 1 - label]
        matches = []
        for candidate_label in candidate_labels:
            image_files, mask_files = indexes[candidate_label]
            if image_key in image_files and image_key in mask_files:
                matches.append(candidate_label)
        if not matches:
            raise FileNotFoundError(
                f"No matched image-mask pair found for cohort case: {image_key}"
            )

        source_label = label if label in matches else matches[0]
        if source_label != label:
            fallback_count += 1
        image_files, mask_files = indexes[source_label]
        dataset_name = Config.DATASET_NAMES[source_label]
        records.append(
            {
                "Case_ID": f"{dataset_name}::{image_key}",
                "Label": label,
                "Source_Dataset": dataset_name,
                "Image": image_files[image_key].name,
                "Mask": mask_files[image_key].name,
                "Image_Path": image_files[image_key],
                "Mask_Path": mask_files[image_key],
            }
        )

    if len(records) != 931:
        raise RuntimeError(f"Expected 931 cohort cases, found {len(records)}")
    if len({record["Case_ID"] for record in records}) != len(records):
        raise RuntimeError("Duplicate U-Net case identifiers were found.")
    print(
        f"Loaded {len(records)} cases from the 931-case cohort; "
        f"{fallback_count} used a non-preferred source directory."
    )
    return records


def split_files():
    records = load_cohort_records()
    labels = [record["Label"] for record in records]
    train_val_records, test_records = train_test_split(
        records,
        test_size=Config.TEST_FRACTION,
        random_state=Config.RANDOM_STATE,
        shuffle=True,
        stratify=labels,
    )
    relative_val_fraction = Config.VAL_FRACTION / (
        Config.TRAIN_FRACTION + Config.VAL_FRACTION
    )
    train_records, val_records = train_test_split(
        train_val_records,
        test_size=relative_val_fraction,
        random_state=Config.RANDOM_STATE,
        shuffle=True,
        stratify=[record["Label"] for record in train_val_records],
    )
    return train_records, val_records, test_records


def save_example_figure(test_rows, output_path, n_cases=3):
    selected = test_rows.sort_values("Dice").head(n_cases)
    panel_size = 360
    header_height = 36
    label_width = 210
    column_titles = ["Chest radiograph", "Expert mask", "U-Net mask", "Error map"]
    canvas_width = label_width + panel_size * 4
    canvas_height = header_height + panel_size * len(selected)
    canvas = Image.new("RGB", (canvas_width, canvas_height), "white")
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.load_default()

    for column, title in enumerate(column_titles):
        x = label_width + column * panel_size + panel_size // 2
        text_width = draw.textlength(title, font=font)
        draw.text((x - text_width / 2, 10), title, fill="black", font=font)

    for row_index, (_, row) in enumerate(selected.iterrows()):
        image = np.asarray(
            Image.open(row["Image_Path"]).convert("L").resize(
                (Config.IMG_SIZE, Config.IMG_SIZE),
                resample=Image.Resampling.BILINEAR,
            )
        )
        reference = np.asarray(
            Image.open(row["Mask_Path"]).convert("L").resize(
                (Config.IMG_SIZE, Config.IMG_SIZE),
                resample=Image.Resampling.NEAREST,
            )
        ) > 0
        prediction = np.asarray(
            Image.open(row["Prediction_Mask_Path"]).convert("L")
        ) > 0

        error_map = np.zeros((*reference.shape, 3), dtype=float)
        true_positive = reference & prediction
        false_positive = ~reference & prediction
        false_negative = reference & ~prediction
        error_map[true_positive] = [0.55, 0.55, 0.55]
        error_map[false_positive] = [0.85, 0.15, 0.15]
        error_map[false_negative] = [0.15, 0.30, 0.85]
        panels = [
            Image.fromarray(image)
            .convert("RGB")
            .resize((panel_size, panel_size), Image.Resampling.BILINEAR),
            Image.fromarray((reference.astype(np.uint8) * 255))
            .convert("RGB")
            .resize((panel_size, panel_size), Image.Resampling.NEAREST),
            Image.fromarray((prediction.astype(np.uint8) * 255))
            .convert("RGB")
            .resize((panel_size, panel_size), Image.Resampling.NEAREST),
            Image.fromarray(
                (error_map.clip(0, 1) * 255).astype(np.uint8)
            ).resize((panel_size, panel_size), Image.Resampling.NEAREST),
        ]
        y = header_height + row_index * panel_size
        label_lines = [
            row.get("Case_ID", row["Image"]),
            f"Dice={row['Dice']:.3f}",
            f"HD95={row['HD95_mm']:.1f} mm",
        ]
        for line_index, line in enumerate(label_lines):
            draw.text((8, y + 12 + line_index * 18), line, fill="black", font=font)
        for column, panel in enumerate(panels):
            canvas.paste(
                panel,
                (label_width + column * panel_size, y),
            )
            draw.rectangle(
                [
                    label_width + column * panel_size,
                    y,
                    label_width + (column + 1) * panel_size - 1,
                    y + panel_size - 1,
                ],
                outline="black",
                width=1,
            )

    canvas.save(output_path)


def export_results(test_rows, history, split_manifest):
    Config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    test_rows.to_csv(
        Config.OUTPUT_DIR / "01_Test_Case_Metrics.csv",
        index=False,
        encoding="utf-8-sig",
    )
    split_manifest.to_csv(
        Config.OUTPUT_DIR / "02_Split_Manifest.csv",
        index=False,
        encoding="utf-8-sig",
    )
    history.to_csv(
        Config.OUTPUT_DIR / "03_Training_History.csv",
        index=False,
        encoding="utf-8-sig",
    )

    low_dice = test_rows.sort_values("Dice").head(15)
    low_dice.to_csv(
        Config.OUTPUT_DIR / "04_Low_Dice_Cases.csv",
        index=False,
        encoding="utf-8-sig",
    )

    summary = {
        "Test_N": int(len(test_rows)),
        "Mean_Dice": float(test_rows["Dice"].mean()),
        "SD_Dice": float(test_rows["Dice"].std(ddof=1)),
        "Median_Dice": float(test_rows["Dice"].median()),
        "Min_Dice": float(test_rows["Dice"].min()),
        "Max_Dice": float(test_rows["Dice"].max()),
        "Mean_HD95_mm": float(test_rows["HD95_mm"].mean()),
        "Median_HD95_mm": float(test_rows["HD95_mm"].median()),
        "Mean_ASSD_mm": float(test_rows["ASSD_mm"].mean()),
        "Median_ASSD_mm": float(test_rows["ASSD_mm"].median()),
        "Dice_Below_0.80_N": int(
            (test_rows["Dice"] < Config.FAILURE_DICE_THRESHOLD).sum()
        ),
        "Empty_Prediction_N": int((test_rows["Prediction_Area_Pixels"] == 0).sum()),
        "Area_Ratio_Failure_N": int(
            test_rows["Failure_Reasons"].str.contains(
                "area_ratio", na=False
            ).sum()
        ),
        "Any_Failure_N": int(test_rows["Failure"].sum()),
        "Any_Failure_Rate": float(test_rows["Failure"].mean()),
        "Mean_Inference_ms": float(test_rows["Inference_ms"].mean()),
        "Median_Inference_ms": float(test_rows["Inference_ms"].median()),
    }
    summary_frame = pd.DataFrame(
        [{"Metric": key, "Value": value} for key, value in summary.items()]
    )

    workbook_path = Config.OUTPUT_DIR / "U-Net_Independent_Test_Metrics.xlsx"
    try:
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            summary_frame.to_excel(writer, sheet_name="Summary", index=False)
            test_rows.drop(
                columns=["Image_Path", "Mask_Path", "Prediction_Mask_Path"],
                errors="ignore",
            ).to_excel(writer, sheet_name="Test_Cases", index=False)
            low_dice.drop(
                columns=["Image_Path", "Mask_Path", "Prediction_Mask_Path"],
                errors="ignore",
            ).to_excel(writer, sheet_name="Low_Dice_Cases", index=False)
            history.to_excel(writer, sheet_name="Training_History", index=False)
            split_manifest.to_excel(writer, sheet_name="Split_Manifest", index=False)
        format_workbook(workbook_path)
    except ModuleNotFoundError:
        workbook_path = None
        print(
            "openpyxl is not installed in this environment; CSV, figure and "
            "JSON outputs are still being written. Run "
            "export_unet_test_results.py in the project Python environment "
            "to generate the Excel workbook."
        )

    save_example_figure(
        test_rows,
        Config.OUTPUT_DIR / "05_Example_Low_Dice_Cases.png",
    )

    config = {
        "cohort_file": str(Config.COHORT_FILE),
        "dataset_dirs": {
            Config.DATASET_NAMES[label]: str(root)
            for label, root in Config.DATASET_DIRS.items()
        },
        "matched_image_mask_pairs": int(len(split_manifest)),
        "train_n": int((split_manifest["Split"] == "Train").sum()),
        "validation_n": int((split_manifest["Split"] == "Validation").sum()),
        "test_n": int((split_manifest["Split"] == "Test").sum()),
        "one_image_per_patient": True,
        "split_level": "image-level, equivalent to patient-level",
        "random_state": Config.RANDOM_STATE,
        "image_size": Config.IMG_SIZE,
        "batch_size": Config.BATCH_SIZE,
        "epochs_requested": Config.EPOCHS,
        "learning_rate": Config.LEARNING_RATE,
        "early_stop_patience": Config.EARLY_STOP_PATIENCE,
        "threshold": 0.5,
        "mm_per_model_pixel": Config.MM_PER_MODEL_PIXEL,
        "failure_criteria": {
            "dice_below": Config.FAILURE_DICE_THRESHOLD,
            "empty_prediction": True,
            "area_ratio_outside": [
                Config.FAILURE_AREA_RATIO_LOW,
                Config.FAILURE_AREA_RATIO_HIGH,
            ],
        },
        "best_model": str(Config.BEST_MODEL_PATH),
        "deployment_models": [str(path) for path in Config.DEPLOY_MODEL_PATHS],
        "summary": summary,
    }
    (Config.OUTPUT_DIR / "06_Config.json").write_text(
        json.dumps(config, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return summary, workbook_path


def promote_checkpoint():
    for target in Config.DEPLOY_MODEL_PATHS:
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(Config.BEST_MODEL_PATH, target)
        print(f"Deployment U-Net updated: {target}")


def format_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = border
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = border
        for column in sheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column]
            width = min(max(max((len(value) for value in values), default=8) + 2, 10), 46)
            sheet.column_dimensions[column[0].column_letter].width = width
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
    workbook.save(path)


def main():
    seed_everything(Config.RANDOM_STATE)
    Config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    Config.MODEL_DIR.mkdir(parents=True, exist_ok=True)

    train_records, val_records, test_records = split_files()
    split_manifest = pd.DataFrame(
        [
            *[{**record, "Split": "Train"} for record in train_records],
            *[{**record, "Split": "Validation"} for record in val_records],
            *[{**record, "Split": "Test"} for record in test_records],
        ]
    )
    split_manifest["Image_Path"] = split_manifest["Image_Path"].astype(str)
    split_manifest["Mask_Path"] = split_manifest["Mask_Path"].astype(str)
    split_manifest.to_csv(
        Config.OUTPUT_DIR / "00_Split_Manifest.csv",
        index=False,
        encoding="utf-8-sig",
    )

    train_dataset = CXRDataset(train_records, img_size=Config.IMG_SIZE)
    val_dataset = CXRDataset(val_records, img_size=Config.IMG_SIZE)
    test_dataset = CXRDataset(test_records, img_size=Config.IMG_SIZE)

    train_loader = DataLoader(
        train_dataset,
        batch_size=Config.BATCH_SIZE,
        shuffle=True,
        num_workers=4,
        pin_memory=torch.cuda.is_available(),
    )
    val_loader = DataLoader(
        val_dataset,
        batch_size=Config.BATCH_SIZE,
        shuffle=False,
        num_workers=4,
        pin_memory=torch.cuda.is_available(),
    )
    model = SimpleUNet().to(Config.DEVICE)
    criterion = DiceBCELoss()
    optimizer = optim.Adam(model.parameters(), lr=Config.LEARNING_RATE)
    scaler = torch.cuda.amp.GradScaler(enabled=Config.DEVICE == "cuda")
    scheduler = optim.lr_scheduler.ReduceLROnPlateau(
        optimizer, mode="max", factor=0.5, patience=3
    )

    history = []
    best_dice = -1.0
    epochs_without_improvement = 0
    print(
        f"Training on {len(train_dataset)} images, validating on {len(val_dataset)}, "
        f"and reserving {len(test_dataset)} for the final independent test."
    )

    for epoch in range(1, Config.EPOCHS + 1):
        model.train()
        train_loss = 0.0
        progress = tqdm(train_loader, desc=f"Epoch {epoch}/{Config.EPOCHS}")
        for images, masks in progress:
            images = images.to(Config.DEVICE)
            masks = masks.to(Config.DEVICE)
            optimizer.zero_grad(set_to_none=True)
            with torch.cuda.amp.autocast(enabled=Config.DEVICE == "cuda"):
                logits = model(images)
                loss = criterion(logits, masks)
            scaler.scale(loss).backward()
            scaler.step(optimizer)
            scaler.update()
            train_loss += loss.item()
            progress.set_postfix(loss=f"{loss.item():.4f}")
        train_loss /= max(len(train_loader), 1)
        val_loss, val_dice = validate(model, val_loader, criterion)
        scheduler.step(val_dice)
        history.append(
            {
                "Epoch": epoch,
                "Train_Loss": train_loss,
                "Validation_Loss": val_loss,
                "Validation_Dice": val_dice,
                "Learning_Rate": optimizer.param_groups[0]["lr"],
            }
        )
        print(
            f"Epoch {epoch}: train loss={train_loss:.4f}, "
            f"validation loss={val_loss:.4f}, validation Dice={val_dice:.4f}"
        )
        if val_dice > best_dice:
            best_dice = val_dice
            epochs_without_improvement = 0
            torch.save(model.state_dict(), Config.BEST_MODEL_PATH)
            print(f"Saved new best validation checkpoint: Dice={best_dice:.4f}")
        else:
            epochs_without_improvement += 1
        if epochs_without_improvement >= Config.EARLY_STOP_PATIENCE:
            print("Early stopping triggered.")
            break

    history_frame = pd.DataFrame(history)
    model.load_state_dict(
        torch.load(Config.BEST_MODEL_PATH, map_location=Config.DEVICE, weights_only=True)
    )
    model.eval()

    prediction_dir = Config.OUTPUT_DIR / "predicted_masks"
    prediction_dir.mkdir(parents=True, exist_ok=True)
    test_rows = []
    with torch.no_grad():
        for record in test_records:
            image_path = record["Image_Path"]
            mask_path = record["Mask_Path"]
            image = Image.open(image_path).convert("L").resize(
                (Config.IMG_SIZE, Config.IMG_SIZE),
                resample=Image.Resampling.BILINEAR,
            )
            image_tensor = transforms.ToTensor()(image).unsqueeze(0).to(Config.DEVICE)
            reference = np.asarray(
                Image.open(mask_path).convert("L").resize(
                    (Config.IMG_SIZE, Config.IMG_SIZE),
                    resample=Image.Resampling.NEAREST,
                )
            ) > 0

            if Config.DEVICE == "cuda":
                torch.cuda.synchronize()
            start = time.perf_counter()
            with torch.cuda.amp.autocast(enabled=Config.DEVICE == "cuda"):
                logits = model(image_tensor)
            if Config.DEVICE == "cuda":
                torch.cuda.synchronize()
            inference_ms = (time.perf_counter() - start) * 1000.0
            prediction = (
                torch.sigmoid(logits).squeeze().cpu().numpy() > 0.5
            )
            prediction_name = record["Case_ID"].replace("::", "__") + ".png"
            prediction_path = prediction_dir / prediction_name
            Image.fromarray((prediction.astype(np.uint8) * 255)).save(
                prediction_path
            )
            metrics = segmentation_metrics(
                reference,
                prediction,
                Config.MM_PER_MODEL_PIXEL,
            )
            test_rows.append(
                {
                    "Case_ID": record["Case_ID"],
                    "Label": record["Label"],
                    "Source_Dataset": record["Source_Dataset"],
                    "Image": image_path.name,
                    "Image_Path": str(image_path),
                    "Mask_Path": str(mask_path),
                    "Prediction_Mask_Path": str(prediction_path),
                    **metrics,
                    "Inference_ms": inference_ms,
                }
            )

    test_rows = pd.DataFrame(test_rows)
    summary, workbook_path = export_results(test_rows, history_frame, split_manifest)
    promote_checkpoint()
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"Results saved to: {Config.OUTPUT_DIR}")
    print(f"Independent-test model saved to: {Config.BEST_MODEL_PATH}")
    print(f"Workbook: {workbook_path or 'not generated in this environment'}")


if __name__ == "__main__":
    main()
