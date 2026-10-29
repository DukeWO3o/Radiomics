from __future__ import annotations

import argparse
import io
import json
import math
import os
import platform
import random
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path

import cv2
import joblib
import numpy as np
import pandas as pd
import pydicom
import SimpleITK as sitk
import torch
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pypinyin import lazy_pinyin
from radiomics import featureextractor
from radiomics import setVerbosity
from sklearn.base import BaseEstimator, TransformerMixin

from unet_model import SimpleUNet


BASE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = BASE_DIR.parent
sys.path.insert(0, str(PROJECT_DIR))
from pipeline_paths import (
    DEFAULT_CLINICAL_931,
    DEFAULT_MODEL_DATA_931,
    DEMO_ROOT,
    IMAGE_ROOT,
)


DEFAULT_DICOM_DIR = IMAGE_ROOT / "RMPP" / "00_DCM"
DEFAULT_TEST_METRICS = (
    BASE_DIR / "unet_validation_output" / "01_Test_Case_Metrics.csv"
)
REFERENCE_WORKBOOK = DEFAULT_MODEL_DATA_931
FALLBACK_CLINICAL_WORKBOOK = DEFAULT_CLINICAL_931
MODELS_DIR = BASE_DIR / "models"
MODEL_CONFIG_PATH = MODELS_DIR / "model_config.json"
UNET_MODEL_PATH = MODELS_DIR / "cxr_unet_best.pth"
RADIOMICS_CONFIG_PATH = BASE_DIR / "config" / "radiomics_config.yaml"
DEFAULT_OUTPUT_DIR = BASE_DIR / "benchmark_output"

CLINICAL_FEATURES = ["age", "CRP", "WBC", "LDH"]
THRESHOLD = 0.5


@dataclass
class Assets:
    seg_model: object
    ml_models: dict
    model_config: dict
    load_times_ms: dict


class StrictPreprocessSelector(BaseEstimator, TransformerMixin):

    def __init__(
        self,
        k=8,
        p_threshold=0.05,
        protected_features=None,
        missing_threshold=0.30,
        variance_threshold=0.0,
        random_state=42,
    ):
        self.k = k
        self.p_threshold = p_threshold
        self.protected_features = protected_features
        self.missing_threshold = missing_threshold
        self.variance_threshold = variance_threshold
        self.random_state = random_state

    def _as_frame(self, x):
        if isinstance(x, pd.DataFrame):
            return x.copy()
        return pd.DataFrame(x, columns=self.feature_names_in_)

    def transform(self, x):
        frame = self._as_frame(x)
        imputed = frame[self.missing_keep_].fillna(self.impute_means_)
        imputed = imputed[self.variance_keep_]
        scaled = self.scaler_.transform(imputed)
        indices = [self.variance_keep_.index(f) for f in self.selected_features_]
        return scaled[:, indices]


setattr(sys.modules["__main__"], "StrictPreprocessSelector", StrictPreprocessSelector)


def should_invert_image(ds, img):
    photometric = str(getattr(ds, "PhotometricInterpretation", "")).strip().upper()
    lut_shape = str(getattr(ds, "PresentationLUTShape", "")).strip().upper()
    if photometric == "MONOCHROME1":
        return True
    if photometric == "MONOCHROME2":
        return False
    if lut_shape == "INVERSE":
        return True
    if lut_shape == "IDENTITY":
        return False
    border = np.concatenate(
        [
            img[:2].ravel(),
            img[-2:].ravel(),
            img[:, :2].ravel(),
            img[:, -2:].ravel(),
        ]
    )
    return float(np.median(border)) > 128


def apply_standardization(dicom_source):
    ds = pydicom.dcmread(dicom_source)
    img = ds.pixel_array.astype(np.float32)
    if "WindowCenter" in ds and "WindowWidth" in ds:
        wc = (
            ds.WindowCenter[0]
            if isinstance(ds.WindowCenter, pydicom.multival.MultiValue)
            else ds.WindowCenter
        )
        ww = (
            ds.WindowWidth[0]
            if isinstance(ds.WindowWidth, pydicom.multival.MultiValue)
            else ds.WindowWidth
        )
        img = np.clip(
            img,
            float(wc) - float(ww) / 2,
            float(wc) + float(ww) / 2,
        )
        img = (
            (img - (float(wc) - float(ww) / 2)) / float(ww) * 255.0
        )
    else:
        img = np.clip(img, np.percentile(img, 1), np.percentile(img, 99))
        span = float(img.max() - img.min())
        img = (img - img.min()) / (span if span > 0 else 1.0) * 255.0
    img = img.astype(np.uint8)

    if "PixelSpacing" in ds:
        spacing = [float(ds.PixelSpacing[0]), float(ds.PixelSpacing[1])]
        factors = [spacing[0] / 0.5, spacing[1] / 0.5]
        new_h = int(img.shape[0] * factors[0])
        new_w = int(img.shape[1] * factors[1])
        if new_h > 0 and new_w > 0:
            img = cv2.resize(
                img,
                (new_w, new_h),
                interpolation=cv2.INTER_AREA,
            )

    if should_invert_image(ds, img):
        img = 255 - img

    target_size = (1024, 1024)
    old_h, old_w = img.shape[:2]
    scale = min(target_size[0] / old_w, target_size[1] / old_h)
    new_w, new_h = int(old_w * scale), int(old_h * scale)
    resized = cv2.resize(
        img,
        (new_w, new_h),
        interpolation=cv2.INTER_AREA,
    )
    canvas = np.zeros(target_size, dtype=np.uint8)
    top = (target_size[1] - new_h) // 2
    left = (target_size[0] - new_w) // 2
    canvas[top : top + new_h, left : left + new_w] = resized
    return canvas


def run_segmentation(img_array, model):
    img_input = cv2.resize(img_array, (512, 512)) / 255.0
    img_tensor = torch.from_numpy(img_input).float().unsqueeze(0).unsqueeze(0)
    with torch.no_grad():
        output = model(img_tensor)
        mask = (
            (torch.sigmoid(output) > 0.5)
            .squeeze()
            .cpu()
            .numpy()
            .astype(np.uint8)
        )
    return cv2.resize(mask, (1024, 1024), interpolation=cv2.INTER_NEAREST)


def create_visual_overlay(img, mask):
    img_color = cv2.cvtColor(img, cv2.COLOR_GRAY2RGB)
    mask_color = np.zeros_like(img_color)
    mask_color[mask == 1] = [0, 255, 0]
    overlay = cv2.addWeighted(img_color, 0.7, mask_color, 0.3, 0)
    contours, _ = cv2.findContours(
        mask,
        cv2.RETR_EXTERNAL,
        cv2.CHAIN_APPROX_SIMPLE,
    )
    cv2.drawContours(overlay, contours, -1, (0, 255, 0), 2)
    return overlay


def load_assets() -> Assets:
    load_times = {}

    started = time.perf_counter()
    seg_model = SimpleUNet(n_channels=1, n_classes=1)
    state = torch.load(str(UNET_MODEL_PATH), map_location="cpu")
    seg_model.load_state_dict(state)
    seg_model.eval()
    load_times["unet_model_load_ms"] = (time.perf_counter() - started) * 1000.0

    with open(MODEL_CONFIG_PATH, "r", encoding="utf-8") as handle:
        model_config = json.load(handle)

    started = time.perf_counter()
    models = {}
    for name in model_config["model_names"]:
        model_path = MODELS_DIR / f"model_{name}.pkl"
        models[name] = joblib.load(str(model_path))
    load_times["tabular_models_load_ms"] = (time.perf_counter() - started) * 1000.0
    load_times["all_assets_load_ms"] = (
        load_times["unet_model_load_ms"]
        + load_times["tabular_models_load_ms"]
    )
    return Assets(seg_model, models, model_config, load_times)


def parse_numeric(value):
    if pd.isna(value):
        return math.nan
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value))
    return float(match.group(0)) if match else math.nan


def load_clinical_inputs(test_names) -> tuple[pd.DataFrame, dict]:
    reference = pd.read_excel(REFERENCE_WORKBOOK, sheet_name="Sheet1")
    reference["Image"] = reference["image"].map(lambda value: Path(str(value)).name)
    reference = reference.drop_duplicates("Image", keep="first").set_index("Image")

    fallback = pd.read_excel(FALLBACK_CLINICAL_WORKBOOK, sheet_name="Sheet1")
    fallback["Image"] = fallback["image"].map(lambda value: Path(str(value)).name)
    fallback = fallback.drop_duplicates("Image", keep="first").set_index("Image")

    medians = {}
    for feature in CLINICAL_FEATURES:
        medians[feature] = float(
            pd.to_numeric(reference[feature], errors="coerce").median()
        )

    rows = []
    imputed_counts = {feature: 0 for feature in CLINICAL_FEATURES}
    for image_name in test_names:
        values = {}
        source = reference if image_name in reference.index else fallback
        if image_name not in source.index:
            raise ValueError(
                f"No clinical inputs were found for anonymized case mapping: {image_name}"
            )
        row = source.loc[image_name]
        for feature in CLINICAL_FEATURES:
            values[feature] = parse_numeric(row[feature])
            if not math.isfinite(values[feature]):
                values[feature] = medians[feature]
                imputed_counts[feature] += 1
        values["Image"] = image_name
        rows.append(values)
    return pd.DataFrame(rows), imputed_counts


def build_dicom_map(dicom_dir: Path) -> dict[str, Path]:
    dicom_files = [
        path
        for path in dicom_dir.iterdir()
        if path.is_file() and path.suffix.lower() == ".dcm"
    ]
    counts = {}
    mapping = {}
    for path in sorted(dicom_files, key=lambda item: item.name):
        pinyin_base = "".join(lazy_pinyin(path.stem)).lower()
        counts[pinyin_base] = counts.get(pinyin_base, 0) + 1
        suffix = "" if counts[pinyin_base] == 1 else f"_{counts[pinyin_base]}"
        mapping[f"{pinyin_base}{suffix}.png"] = path
    return mapping


def wilson_interval(successes: int, total: int, z: float = 1.959963984540054):
    if total == 0:
        return math.nan, math.nan
    proportion = successes / total
    denominator = 1 + z * z / total
    center = (proportion + z * z / (2 * total)) / denominator
    half_width = (
        z
        * math.sqrt(
            proportion * (1 - proportion) / total + z * z / (4 * total * total)
        )
        / denominator
    )
    return center - half_width, center + half_width


def percentile_summary(values) -> dict:
    values = np.asarray([value for value in values if math.isfinite(value)], dtype=float)
    if not len(values):
        return {
            "N": 0,
            "Mean_ms": math.nan,
            "SD_ms": math.nan,
            "Median_ms": math.nan,
            "Q1_ms": math.nan,
            "Q3_ms": math.nan,
            "P95_ms": math.nan,
            "Min_ms": math.nan,
            "Max_ms": math.nan,
        }
    return {
        "N": int(len(values)),
        "Mean_ms": float(np.mean(values)),
        "SD_ms": float(np.std(values, ddof=1)) if len(values) > 1 else math.nan,
        "Median_ms": float(np.median(values)),
        "Q1_ms": float(np.percentile(values, 25)),
        "Q3_ms": float(np.percentile(values, 75)),
        "P95_ms": float(np.percentile(values, 95)),
        "Min_ms": float(np.min(values)),
        "Max_ms": float(np.max(values)),
    }


def classify_failure(stage: str, exception: Exception) -> str:
    name = type(exception).__name__
    if stage == "dicom_read_and_standardization":
        return f"DICOM_or_standardization_error:{name}"
    if stage == "unet_preprocess_inference_mask_resize":
        return f"U-Net_error:{name}"
    if stage.startswith("radiomics"):
        return f"Radiomics_error:{name}"
    if stage == "feature_alignment":
        return f"Feature_alignment_error:{name}"
    if stage == "model_prediction_all9":
        return f"Model_prediction_error:{name}"
    return f"{stage}:{name}"


def run_single_case(
    dicom_bytes: bytes,
    clinical_row,
    assets: Assets,
) -> tuple[dict, dict]:
    timings = {}
    values = {}
    stage = "initialization"
    try:
        core_started = time.perf_counter()

        stage = "dicom_read_and_standardization"
        started = time.perf_counter()
        std_img = apply_standardization(io.BytesIO(dicom_bytes))
        timings[stage] = (time.perf_counter() - started) * 1000.0

        stage = "unet_preprocess_inference_mask_resize"
        started = time.perf_counter()
        mask = run_segmentation(std_img, assets.seg_model)
        timings[stage] = (time.perf_counter() - started) * 1000.0
        if int(mask.sum()) == 0:
            raise ValueError("Empty U-Net mask")

        stage = "overlay_generation_auxiliary"
        started = time.perf_counter()
        _ = create_visual_overlay(std_img, mask)
        timings[stage] = (time.perf_counter() - started) * 1000.0

        stage = "radiomics_extractor_init"
        started = time.perf_counter()
        sitk_img = sitk.GetImageFromArray(std_img)
        sitk_mask = sitk.GetImageFromArray(mask)
        extractor = featureextractor.RadiomicsFeatureExtractor(
            str(RADIOMICS_CONFIG_PATH)
        )
        timings[stage] = (time.perf_counter() - started) * 1000.0

        stage = "radiomics_extraction"
        started = time.perf_counter()
        features = extractor.execute(sitk_img, sitk_mask)
        timings[stage] = (time.perf_counter() - started) * 1000.0

        stage = "feature_alignment"
        started = time.perf_counter()
        raw_feature_dict = {
            key.replace("original_", ""): float(value)
            for key, value in features.items()
            if "diagnostics" not in key
        }
        clinical_inputs = {
            feature: float(clinical_row[feature]) for feature in CLINICAL_FEATURES
        }
        feature_order = assets.model_config["feature_order"]
        full_row = {
            column: clinical_inputs.get(
                column,
                raw_feature_dict.get(column, np.nan),
            )
            for column in feature_order
        }
        input_df = pd.DataFrame([full_row], columns=feature_order).astype(float)
        timings[stage] = (time.perf_counter() - started) * 1000.0

        stage = "model_prediction_all9"
        started = time.perf_counter()
        results = {
            name: float(model.predict_proba(input_df)[0, 1])
            for name, model in assets.ml_models.items()
        }
        timings[stage] = (time.perf_counter() - started) * 1000.0

        stage = "result_validation"
        started = time.perf_counter()
        best_model_name = assets.model_config.get("best_model", "SVM")
        probability = results.get(best_model_name, results.get("SVM"))
        if probability is None or not math.isfinite(probability):
            raise ValueError("Final probability is missing or non-finite")
        if not 0.0 <= probability <= 1.0:
            raise ValueError("Final probability is outside [0, 1]")
        timings[stage] = (time.perf_counter() - started) * 1000.0

        timings["app_compute_total"] = (
            time.perf_counter() - core_started
        ) * 1000.0
        timings["core_total_excluding_overlay"] = (
            timings["app_compute_total"]
            - timings["overlay_generation_auxiliary"]
        )
        values = {
            "Probability": probability,
            "Predicted_Class": int(probability >= THRESHOLD),
            "Technical_Success": True,
            "Failure_Stage": "",
            "Failure_Type": "",
        }
    except Exception as exception:
        values = {
            "Probability": math.nan,
            "Predicted_Class": np.nan,
            "Technical_Success": False,
            "Failure_Stage": stage,
            "Failure_Type": classify_failure(stage, exception),
        }
        timings["app_compute_total"] = (
            time.perf_counter() - core_started
        ) * 1000.0
        if "overlay_generation_auxiliary" in timings:
            timings["core_total_excluding_overlay"] = (
                timings["app_compute_total"]
                - timings["overlay_generation_auxiliary"]
            )
        else:
            timings["core_total_excluding_overlay"] = timings["app_compute_total"]
    return values, timings


def warm_up(assets: Assets, warmup_runs: int) -> list[dict]:
    demo_meta = {
        "rmpp_case_positive.dcm": (4.92, 9.28, 11.15, 391.0),
        "rmpp_case_negative.dcm": (3.42, 0.29, 7.85, 299.0),
    }
    rows = []
    demo_paths = sorted(DEMO_ROOT.glob("*.dcm"))
    for index in range(warmup_runs):
        path = demo_paths[index % len(demo_paths)]
        age, crp, wbc, ldh = demo_meta[path.name]
        values, timings = run_single_case(
            path.read_bytes(),
            pd.Series({"age": age, "CRP": crp, "WBC": wbc, "LDH": ldh}),
            assets,
        )
        rows.append(
            {
                "Warmup": index + 1,
                "Technical_Success": values["Technical_Success"],
                "Total_ms": timings["app_compute_total"],
            }
        )
    return rows


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                        wrap_text=True,
                    )
                else:
                    cell.font = body_font
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
        for column in worksheet.columns:
            maximum = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column
            )
            worksheet.column_dimensions[column[0].column_letter].width = min(
                max(maximum + 2, 12),
                54,
            )
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
    workbook.save(path)


def write_outputs(
    output_dir: Path,
    raw_results: pd.DataFrame,
    case_summary: pd.DataFrame,
    stage_summary: pd.DataFrame,
    overall_summary: dict,
    asset_load_times: dict,
    warmup_results: list[dict],
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = output_dir / "End_to_End_Timing_Results.xlsx"
    readme = pd.DataFrame(
        [
            (
                "Purpose",
                "Warm-run timing of the non-Streamlit computation path used by "
                "06_RMPP_Prediction_App_CV/app.py.",
            ),
            (
                "Primary endpoint",
                "DICOM bytes available in memory to a valid final SVM probability.",
            ),
            (
                "Excluded from primary endpoint",
                "Browser upload, disk read, application startup, Streamlit page "
                "rendering, and user idle time.",
            ),
            (
                "app_compute_total",
                "Includes DICOM standardization, U-Net CPU inference, mask resize, "
                "overlay generation, PyRadiomics, feature alignment, and nine-model "
                "prediction.",
            ),
            (
                "core_total_excluding_overlay",
                "Same as app_compute_total but excludes the optional visual overlay.",
            ),
            (
                "Repeats",
                "Each case was run multiple times in randomized order; per-case "
                "medians are used for across-case summaries.",
            ),
            (
                "Technical success",
                "All stages completed and the final probability was finite and in [0,1].",
            ),
            (
                "Privacy",
                "Only UNET-### identifiers and timing/result fields are exported; "
                "source filenames and paths are not written.",
            ),
            (
                "Clinical input fallback",
                "Missing clinical values are replaced by reference-cohort medians "
                "only to exercise the technical pipeline. Benchmark probabilities "
                "are not a new clinical performance analysis.",
            ),
        ],
        columns=["Item", "Description"],
    )
    summary_frame = pd.DataFrame(
        [{"Metric": key, "Value": value} for key, value in overall_summary.items()]
    )
    load_frame = pd.DataFrame(
        [{"Metric": key, "Value": value} for key, value in asset_load_times.items()]
    )
    warmup_frame = pd.DataFrame(warmup_results)
    failures = raw_results.loc[
        ~raw_results["Technical_Success"].astype(bool),
        [
            "Case_ID",
            "Repetition",
            "Failure_Stage",
            "Failure_Type",
            "app_compute_total_ms",
        ],
    ].copy()

    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        summary_frame.to_excel(writer, sheet_name="Summary", index=False)
        stage_summary.to_excel(writer, sheet_name="Stage_Summary", index=False)
        case_summary.to_excel(writer, sheet_name="Case_Summary", index=False)
        raw_results.to_excel(writer, sheet_name="All_Repetitions", index=False)
        failures.to_excel(writer, sheet_name="Failures", index=False)
        load_frame.to_excel(writer, sheet_name="Asset_Load", index=False)
        warmup_frame.to_excel(writer, sheet_name="Warmup", index=False)
        pd.DataFrame([platform_info()]).to_excel(
            writer,
            sheet_name="Environment",
            index=False,
        )
    style_workbook(workbook_path)

    raw_results.to_csv(
        output_dir / "all_repetitions.csv",
        index=False,
        encoding="utf-8-sig",
    )
    case_summary.to_csv(
        output_dir / "case_summary.csv",
        index=False,
        encoding="utf-8-sig",
    )
    stage_summary.to_csv(
        output_dir / "stage_summary.csv",
        index=False,
        encoding="utf-8-sig",
    )
    (output_dir / "summary.json").write_text(
        json.dumps(overall_summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return workbook_path


def platform_info() -> dict:
    import sklearn

    return {
        "Operating_System": platform.platform(),
        "Machine": platform.machine(),
        "Processor": platform.processor(),
        "Logical_CPU_Count": os.cpu_count(),
        "Python": platform.python_version(),
        "PyTorch": torch.__version__,
        "scikit_learn": sklearn.__version__,
        "PyRadiomics": getattr(featureextractor, "__version__", "unknown"),
        "SimpleITK": sitk.Version_VersionString(),
        "pydicom": pydicom.__version__,
        "Inference_Device": "cpu",
    }


def find_test_metrics(
    dicom_map: dict[str, Path],
) -> tuple[pd.DataFrame, dict[str, str]]:
    metrics = pd.read_csv(DEFAULT_TEST_METRICS, encoding="utf-8-sig")
    names = metrics["Image"].astype(str).map(lambda value: Path(value).name)
    missing_dicoms = [name for name in names if name not in dicom_map]
    if missing_dicoms:
        raise ValueError(
            f"{len(missing_dicoms)} test cases could not be mapped to a DICOM file."
        )
    case_ids = {
        name: f"UNET-{index:03d}"
        for index, name in enumerate(sorted(set(names)), start=1)
    }
    metrics = metrics.copy()
    metrics["Image"] = names
    metrics["Case_ID"] = metrics["Image"].map(case_ids)
    metrics["Dicom_Path"] = metrics["Image"].map(dicom_map)
    return metrics, case_ids


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Benchmark the non-UI DICOM-to-probability path used by the RMPP app."
        )
    )
    parser.add_argument(
        "--dicom-dir",
        type=Path,
        default=DEFAULT_DICOM_DIR,
    )
    parser.add_argument("--repeat", type=int, default=3)
    parser.add_argument("--warmup", type=int, default=2)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Limit the number of cases for a smoke test.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
    )
    args = parser.parse_args()
    if args.repeat < 1:
        raise ValueError("--repeat must be at least 1")

    setVerbosity(40)
    random.seed(args.seed)
    np.random.seed(args.seed)
    torch.manual_seed(args.seed)

    dicom_map = build_dicom_map(args.dicom_dir)
    metrics, _ = find_test_metrics(dicom_map)
    benchmark = metrics
    if args.limit is not None:
        benchmark = benchmark.head(args.limit).copy()

    clinical, clinical_imputations = load_clinical_inputs(
        benchmark["Image"].tolist()
    )
    benchmark = benchmark.merge(
        clinical,
        on="Image",
        how="left",
        validate="one_to_one",
    )

    assets = load_assets()
    warmup_results = warm_up(assets, args.warmup)

    tasks = [
        (row.Case_ID, row.Image, row.Dicom_Path, repetition)
        for row in benchmark.itertuples(index=False)
        for repetition in range(1, args.repeat + 1)
    ]
    random.shuffle(tasks)
    clinical_lookup = benchmark.set_index("Image")

    rows = []
    for task_index, (case_id, image_name, dicom_path, repetition) in enumerate(
        tasks,
        start=1,
    ):
        preload_started = time.perf_counter()
        dicom_bytes = Path(dicom_path).read_bytes()
        preload_ms = (time.perf_counter() - preload_started) * 1000.0

        values, timings = run_single_case(
            dicom_bytes,
            clinical_lookup.loc[image_name],
            assets,
        )
        rows.append(
            {
                "Case_ID": case_id,
                "Repetition": repetition,
                "Preload_ms_excluded": preload_ms,
                **values,
                **{f"{key}_ms": value for key, value in timings.items()},
            }
        )
        if task_index % 10 == 0 or task_index == len(tasks):
            print(f"Benchmarked {task_index}/{len(tasks)} runs")

    raw_results = pd.DataFrame(rows)

    component_stages = [
        "dicom_read_and_standardization",
        "unet_preprocess_inference_mask_resize",
        "overlay_generation_auxiliary",
        "radiomics_extractor_init",
        "radiomics_extraction",
        "feature_alignment",
        "model_prediction_all9",
        "result_validation",
    ]
    stage_columns = [f"{stage}_ms" for stage in component_stages]
    successful = raw_results[raw_results["Technical_Success"]].copy()
    repetitions = (
        successful.groupby("Case_ID", as_index=False)
        .size()
        .rename(columns={"size": "Successful_Repetitions"})
    )
    case_medians = successful.groupby("Case_ID", as_index=False)[
        stage_columns
        + ["app_compute_total_ms", "core_total_excluding_overlay_ms"]
    ].median()
    probabilities = successful.groupby("Case_ID", as_index=False)[
        "Probability"
    ].median()
    case_summary = (
        repetitions.merge(case_medians, on="Case_ID", how="left")
        .merge(probabilities, on="Case_ID", how="left")
    )

    stage_rows = []
    for stage, column in zip(component_stages, stage_columns):
        stage_rows.append(
            {
                "Stage": stage,
                **percentile_summary(case_summary[column]),
            }
        )
    stage_summary = pd.DataFrame(stage_rows)
    total_median_sum = stage_summary["Median_ms"].sum()
    stage_summary["Share_of_app_compute_percent"] = (
        stage_summary["Median_ms"] / total_median_sum * 100.0
    )

    case_ids = benchmark["Case_ID"].tolist()
    case_success = {
        case_id: bool(
            raw_results.loc[
                raw_results["Case_ID"] == case_id,
                "Technical_Success",
            ].any()
        )
        for case_id in case_ids
    }
    case_all_success = {
        case_id: bool(
            raw_results.loc[
                raw_results["Case_ID"] == case_id,
                "Technical_Success",
            ].all()
        )
        for case_id in case_ids
    }
    successful_cases = sum(case_success.values())
    all_successful_cases = sum(case_all_success.values())
    success_ci_low, success_ci_high = wilson_interval(
        successful_cases,
        len(case_ids),
    )
    per_case_totals = successful.groupby("Case_ID")[
        "app_compute_total_ms"
    ].median()
    core_totals = successful.groupby("Case_ID")[
        "core_total_excluding_overlay_ms"
    ].median()
    total_stats = percentile_summary(per_case_totals)
    core_stats = percentile_summary(core_totals)

    overall_summary = {
        "Planned_Cases": int(len(case_ids)),
        "Repeats_Per_Case": int(args.repeat),
        "Total_Attempts": int(len(raw_results)),
        "Successful_Attempts": int(raw_results["Technical_Success"].sum()),
        "Attempt_Success_Rate": float(raw_results["Technical_Success"].mean()),
        "Cases_With_At_Least_One_Success": int(successful_cases),
        "Case_Level_Success_Rate": float(successful_cases / len(case_ids)),
        "Case_Level_Success_Rate_95CI_Low": success_ci_low,
        "Case_Level_Success_Rate_95CI_High": success_ci_high,
        "Cases_With_All_Repeats_Successful": int(all_successful_cases),
        "All_Repeats_Success_Rate": float(all_successful_cases / len(case_ids)),
        "Median_app_compute_total_ms": total_stats["Median_ms"],
        "Q1_app_compute_total_ms": total_stats["Q1_ms"],
        "Q3_app_compute_total_ms": total_stats["Q3_ms"],
        "P95_app_compute_total_ms": total_stats["P95_ms"],
        "Mean_app_compute_total_ms": total_stats["Mean_ms"],
        "SD_app_compute_total_ms": total_stats["SD_ms"],
        "Median_core_total_excluding_overlay_ms": core_stats["Median_ms"],
        "Q1_core_total_excluding_overlay_ms": core_stats["Q1_ms"],
        "Q3_core_total_excluding_overlay_ms": core_stats["Q3_ms"],
        "P95_core_total_excluding_overlay_ms": core_stats["P95_ms"],
        "Warmup_Runs": int(args.warmup),
        "Random_Seed": int(args.seed),
        "Clinical_Features_Imputed_For_Timing": int(
            sum(clinical_imputations.values())
        ),
        "Clinical_Imputation_Counts": json.dumps(
            clinical_imputations,
            ensure_ascii=False,
        ),
    }

    workbook_path = write_outputs(
        args.output_dir,
        raw_results,
        case_summary,
        stage_summary,
        overall_summary,
        assets.load_times_ms,
        warmup_results,
    )
    print(json.dumps(overall_summary, ensure_ascii=False, indent=2))
    print(f"Results: {workbook_path}")


if __name__ == "__main__":
    main()
