from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "unet_validation_output"
METRICS_CSV = OUTPUT_DIR / "01_Test_Case_Metrics.csv"
HISTORY_CSV = OUTPUT_DIR / "03_Training_History.csv"
SPLIT_CSV = OUTPUT_DIR / "02_Split_Manifest.csv"
MODEL_PATH = BASE_DIR / "models" / "cxr_unet_best_independent.pth"
WORKBOOK_PATH = OUTPUT_DIR / "U-Net_Independent_Test_Metrics.xlsx"
FIGURE_PATH = OUTPUT_DIR / "05_Example_Low_Dice_Cases.png"
CONFIG_PATH = OUTPUT_DIR / "06_Config.json"

IMAGE_SIZE = 512
MM_PER_MODEL_PIXEL = 1.0
FAILURE_DICE_THRESHOLD = 0.80


def save_example_figure(test_rows, n_cases=3):
    selected = test_rows.sort_values("Dice").head(n_cases)
    panel_size = 360
    header_height = 36
    label_width = 220
    column_titles = ["Chest radiograph", "Expert mask", "U-Net mask", "Error map"]
    canvas = Image.new(
        "RGB",
        (label_width + panel_size * 4, header_height + panel_size * len(selected)),
        "white",
    )
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.load_default()

    for column, title in enumerate(column_titles):
        center_x = label_width + column * panel_size + panel_size // 2
        width = draw.textlength(title, font=font)
        draw.text((center_x - width / 2, 10), title, fill="black", font=font)

    for row_index, (_, row) in enumerate(selected.iterrows()):
        image = Image.open(row["Image_Path"]).convert("L").resize(
            (IMAGE_SIZE, IMAGE_SIZE),
            resample=Image.Resampling.BILINEAR,
        )
        reference = np.asarray(
            Image.open(row["Mask_Path"])
            .convert("L")
            .resize((IMAGE_SIZE, IMAGE_SIZE), resample=Image.Resampling.NEAREST)
        ) > 0
        prediction = np.asarray(
            Image.open(row["Prediction_Mask_Path"]).convert("L")
        ) > 0

        error_map = np.zeros((IMAGE_SIZE, IMAGE_SIZE, 3), dtype=float)
        error_map[reference & prediction] = [0.55, 0.55, 0.55]
        error_map[~reference & prediction] = [0.85, 0.15, 0.15]
        error_map[reference & ~prediction] = [0.15, 0.30, 0.85]

        panels = [
            image.convert("RGB").resize(
                (panel_size, panel_size), resample=Image.Resampling.BILINEAR
            ),
            Image.fromarray((reference.astype(np.uint8) * 255))
            .convert("RGB")
            .resize((panel_size, panel_size), resample=Image.Resampling.NEAREST),
            Image.fromarray((prediction.astype(np.uint8) * 255))
            .convert("RGB")
            .resize((panel_size, panel_size), resample=Image.Resampling.NEAREST),
            Image.fromarray((error_map * 255).clip(0, 255).astype(np.uint8)).resize(
                (panel_size, panel_size), resample=Image.Resampling.NEAREST
            ),
        ]
        row_y = header_height + row_index * panel_size
        labels = [
            row["Image"],
            f"Dice={row['Dice']:.3f}",
            f"HD95={row['HD95_mm']:.1f} mm",
            f"ASSD={row['ASSD_mm']:.1f} mm",
        ]
        for line_index, label in enumerate(labels):
            draw.text(
                (8, row_y + 12 + line_index * 18),
                label,
                fill="black",
                font=font,
            )
        for column, panel in enumerate(panels):
            x = label_width + column * panel_size
            canvas.paste(panel, (x, row_y))
            draw.rectangle(
                [x, row_y, x + panel_size - 1, row_y + panel_size - 1],
                outline="black",
                width=1,
            )

    canvas.save(FIGURE_PATH)


def format_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = border
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = border
        for column in sheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column]
            width = min(max(max((len(value) for value in values), default=8) + 2, 10), 46)
            sheet.column_dimensions[column[0].column_letter].width = width
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
    workbook.save(path)


def main():
    metrics = pd.read_csv(METRICS_CSV)
    history = pd.read_csv(HISTORY_CSV)
    split_manifest = pd.read_csv(SPLIT_CSV)
    low_dice = metrics.sort_values("Dice").head(15)

    summary_values = {
        "Test_N": int(len(metrics)),
        "Train_N": int((split_manifest["Split"] == "Train").sum()),
        "Validation_N": int((split_manifest["Split"] == "Validation").sum()),
        "Best_Validation_Dice": float(history["Validation_Dice"].max()),
        "Mean_Dice": float(metrics["Dice"].mean()),
        "SD_Dice": float(metrics["Dice"].std(ddof=1)),
        "Median_Dice": float(metrics["Dice"].median()),
        "Min_Dice": float(metrics["Dice"].min()),
        "Max_Dice": float(metrics["Dice"].max()),
        "Mean_HD95_mm": float(metrics["HD95_mm"].mean()),
        "Median_HD95_mm": float(metrics["HD95_mm"].median()),
        "Mean_ASSD_mm": float(metrics["ASSD_mm"].mean()),
        "Median_ASSD_mm": float(metrics["ASSD_mm"].median()),
        "Dice_Below_0.80_N": int((metrics["Dice"] < FAILURE_DICE_THRESHOLD).sum()),
        "Empty_Prediction_N": int((metrics["Prediction_Area_Pixels"] == 0).sum()),
        "Area_Ratio_Failure_N": int(
            metrics["Failure_Reasons"].fillna("").str.contains("area_ratio").sum()
        ),
        "Any_Failure_N": int(metrics["Failure"].sum()),
        "Any_Failure_Rate": float(metrics["Failure"].mean()),
        "Mean_Inference_ms": float(metrics["Inference_ms"].mean()),
        "Median_Inference_ms": float(metrics["Inference_ms"].median()),
    }
    summary = pd.DataFrame(
        [{"Metric": key, "Value": value} for key, value in summary_values.items()]
    )

    columns_to_drop = ["Image_Path", "Mask_Path", "Prediction_Mask_Path"]
    with pd.ExcelWriter(WORKBOOK_PATH, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        metrics.drop(columns=columns_to_drop, errors="ignore").to_excel(
            writer, sheet_name="Test_Cases", index=False
        )
        low_dice.drop(columns=columns_to_drop, errors="ignore").to_excel(
            writer, sheet_name="Low_Dice_Cases", index=False
        )
        history.to_excel(writer, sheet_name="Training_History", index=False)
        split_manifest.to_excel(writer, sheet_name="Split_Manifest", index=False)
    format_workbook(WORKBOOK_PATH)

    save_example_figure(metrics)

    config = {
        "one_image_per_patient": True,
        "split_level": "image-level, equivalent to patient-level",
        "random_state": 42,
        "image_size": IMAGE_SIZE,
        "threshold": 0.5,
        "mm_per_model_pixel": MM_PER_MODEL_PIXEL,
        "failure_criteria": {
            "dice_below": FAILURE_DICE_THRESHOLD,
            "empty_prediction": True,
            "area_ratio_outside": [0.50, 2.00],
        },
        "best_model": str(MODEL_PATH),
        "best_model_size_bytes": MODEL_PATH.stat().st_size,
        "summary": summary_values,
    }
    CONFIG_PATH.write_text(
        json.dumps(config, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(json.dumps(summary_values, ensure_ascii=False, indent=2))
    print(f"Workbook: {WORKBOOK_PATH}")
    print(f"Example figure: {FIGURE_PATH}")


if __name__ == "__main__":
    main()
