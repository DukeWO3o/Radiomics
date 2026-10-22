from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path(
    r"E:\Softwares\PycharmProject\Radiomics\Pediatric Radiology"
    r"\Cover_Letter_Pediatric_Radiology.docx"
)

TITLE = (
    "Prediction of refractory Mycoplasma pneumoniae pneumonia in children "
    "using chest radiography-based radiomics and machine learning"
)


def set_run_font(run, *, size: float = 11, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")


def add_paragraph(
    doc: Document,
    text: str = "",
    *,
    bold: bool = False,
    space_after: float = 5,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)


def build_cover_letter() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")

    add_paragraph(doc, "Bofan Wang", bold=True, space_after=0)
    add_paragraph(
        doc,
        "Children's Hospital of Nanjing Medical University",
        space_after=0,
    )
    add_paragraph(doc, "Nanjing 211100, Jiangsu, China", space_after=12)

    add_paragraph(doc, "September 15, 2026", space_after=12)
    add_paragraph(doc, "Editor-in-Chief", space_after=0)
    add_paragraph(doc, "Pediatric Radiology", space_after=12)
    add_paragraph(
        doc,
        f"Re: Submission of an original Research article: \"{TITLE}\"",
        bold=True,
        space_after=12,
    )

    add_paragraph(doc, "Dear Editor-in-Chief,", space_after=8)
    add_paragraph(
        doc,
        (
            "On behalf of all co-authors, I am pleased to submit our original "
            "research article for consideration as a Research article in "
            "Pediatric Radiology."
        ),
    )
    add_paragraph(
        doc,
        (
            "Refractory Mycoplasma pneumoniae pneumonia (RMPP) is an important "
            "clinical problem in pediatric respiratory medicine because delayed "
            "recognition can lead to prolonged illness and complications such as "
            "pleural effusion, necrotizing pneumonia, and bronchiolitis obliterans. "
            "Computed tomography may provide detailed anatomical information, but its "
            "radiation burden limits repeated use in children. Our study therefore "
            "evaluates whether routinely obtained chest radiographs can support "
            "earlier, low-radiation risk stratification."
        ),
    )
    add_paragraph(
        doc,
        (
            "We retrospectively studied a large cohort of children with Mycoplasma "
            "pneumoniae pneumonia. We assessed radiomics reproducibility, performed "
            "feature selection within each cross-validation fold, and compared "
            "multiple machine learning models using discrimination, calibration, "
            "decision curve analysis, and Shapley additive explanations. The "
            "LightGBM model showed moderate discriminative performance and clinical "
            "utility in internal validation. We also developed a web-based interface "
            "that accepts a DICOM chest radiograph and routine clinical variables and "
            "returns an individualized RMPP risk estimate."
        ),
    )
    add_paragraph(
        doc,
        (
            "We believe this work fits the scope of Pediatric Radiology because it "
            "addresses a common pediatric chest disease, uses an accessible imaging "
            "modality already embedded in pediatric practice, and links quantitative "
            "imaging analysis to a practical clinical workflow. The findings are "
            "relevant to pediatric radiologists, respiratory clinicians, and teams "
            "developing imaging-based decision-support tools. The manuscript is "
            "accompanied by figures, tables, and Online Resources. The main manuscript "
            "and accompanying files have been anonymized for double-blind review; "
            "author information and declarations are provided on the separate Title "
            "Page."
        ),
    )
    add_paragraph(
        doc,
        (
            "This manuscript is original, has not been published previously, and is "
            "not currently under consideration by another journal. All authors have "
            "reviewed the manuscript, approved its submission, and agree to be "
            "accountable for the work. The authors declare no competing interests. "
            "This work was supported by the National Natural Science Foundation of "
            "China (grant no. 82500006). The study was conducted in accordance with "
            "the Declaration of Helsinki, and the ethics approval and consent "
            "statements are provided on the Title Page. Thank you for your time and "
            "consideration."
        ),
    )

    add_paragraph(doc, "Sincerely,", space_after=12)
    add_paragraph(doc, "Bofan Wang", bold=True, space_after=0)
    add_paragraph(doc, "On behalf of all authors", space_after=12)
    add_paragraph(doc, "Corresponding authors:", bold=True, space_after=0)
    add_paragraph(
        doc,
        "Deyu Zhao, zhaodeyu988@126.com",
        space_after=0,
    )
    add_paragraph(
        doc,
        "Xia Huang, huang_xia01@163.com",
        space_after=0,
    )

    doc.core_properties.author = "Bofan Wang"
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "Cover Letter - Pediatric Radiology"
    doc.core_properties.subject = (
        "Submission cover letter for the manuscript submitted to Pediatric Radiology"
    )
    doc.core_properties.comments = ""
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_cover_letter()
