# -*- coding: utf-8 -*-
"""
Fixed-parameter 5x5 repeated stratified cross-validation for RMPP models.

This script reuses the fold-internal preprocessing, univariate filtering,
and mRMR implementation from I.refactored_pipeline_strict_cv.py. It does
not re-tune hyperparameters. For each repeated CV fold:

    missing-rate filter -> mean imputation -> raw-scale constant filter
    -> Z-score -> Mann-Whitney U + FDR -> mRMR -> final classifier

Only the original 744-case training set is resampled. The frozen 187-case
internal validation set is never used in this analysis.

Example:
    python I-5.repeated_5x5_cv_strict.py --modality combined
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import sys
import warnings
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from sklearn.base import clone
from sklearn.metrics import (
    accuracy_score,
    brier_score_loss,
    confusion_matrix,
    recall_score,
    roc_auc_score,
)
from sklearn.model_selection import RepeatedStratifiedKFold, train_test_split


warnings.filterwarnings("ignore")
plt.rcParams["font.sans-serif"] = ["SimHei"]
plt.rcParams["axes.unicode_minus"] = False

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
BASE_PIPELINE = SCRIPT_DIR / "I.refactored_pipeline_strict_cv.py"
STRICT_ROOT = PROJECT_ROOT / "06_Modelling" / "06e_StrictCV_折内预处理"
OUTPUT_LABELS = {
    "combined": "联合模型",
    "clinical": "仅临床变量",
    "radiomics": "仅影像组学",
}
RANDOM_STATE = 42
N_SPLITS = 5
N_REPEATS = 5

# Final parameters from the completed strict-CV tuning run. These values are
# fixed for all 25 repeated-CV folds; no search or selection is performed here.
FINAL_PARAMS = {
    "LogisticRegression": {
        "C": 0.1,
        "penalty": "l2",
        "solver": "liblinear",
    },
    "DecisionTree": {
        "max_depth": 3,
        "min_samples_leaf": 10,
        "min_samples_split": 40,
        "max_features": None,
        "ccp_alpha": 0.0,
    },
    "RandomForest": {
        "n_estimators": 100,
        "max_depth": 3,
        "min_samples_leaf": 10,
        "min_samples_split": 30,
        "max_features": 0.5,
        "bootstrap": True,
    },
    "MLP": {
        "hidden_layer_sizes": [128, 64],
        "alpha": 0.01,
        "activation": "relu",
    },
    "SVM": {
        "C": 2,
        "gamma": 0.025,
        "kernel": "rbf",
    },
    "XGBoost": {
        "n_estimators": 30,
        "max_depth": 2,
        "learning_rate": 0.05,
        "subsample": 0.8,
        "colsample_bytree": 0.5,
        "min_child_weight": 10,
        "reg_alpha": 2,
        "reg_lambda": 10,
        "scale_pos_weight": 0.67,
    },
    "LightGBM": {
        "n_estimators": 50,
        "max_depth": 3,
        "learning_rate": 0.03,
        "num_leaves": 7,
        "min_child_samples": 100,
        "reg_lambda": 0.6,
    },
    "GaussianNB": {
        "var_smoothing": 1e-12,
    },
    "KNN": {
        "n_neighbors": 41,
        "weights": "uniform",
        "metric": "euclidean",
    },
}


def load_base_pipeline():
    spec = importlib.util.spec_from_file_location(
        "strict_cv_base_pipeline", BASE_PIPELINE
    )
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load base pipeline: {BASE_PIPELINE}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def specificity_score(y_true, y_pred):
    tn, fp, _, _ = confusion_matrix(y_true, y_pred, labels=[0, 1]).ravel()
    return tn / (tn + fp) if (tn + fp) else np.nan


def positive_rate_metrics(y_true, y_prob):
    y_pred = (y_prob >= 0.5).astype(int)
    return {
        "AUC": roc_auc_score(y_true, y_prob),
        "Brier": brier_score_loss(y_true, y_prob),
        "Accuracy": accuracy_score(y_true, y_pred),
        "Recall": recall_score(y_true, y_pred, zero_division=0),
        "Specificity": specificity_score(y_true, y_pred),
    }


def summarize_model_auc(fold_frame):
    repeat_means = (
        fold_frame.groupby("Repeat", as_index=False)["AUC"].mean()
        .rename(columns={"AUC": "Repeat_Mean_AUC"})
    )
    fold_auc = fold_frame["AUC"].to_numpy(dtype=float)
    repeat_auc = repeat_means["Repeat_Mean_AUC"].to_numpy(dtype=float)
    return {
        "Pooled_Fold_Mean_AUC": float(np.mean(fold_auc)),
        "Pooled_Fold_SD": float(np.std(fold_auc, ddof=1)),
        "Pooled_Fold_Median_AUC": float(np.median(fold_auc)),
        "Pooled_Fold_Min_AUC": float(np.min(fold_auc)),
        "Pooled_Fold_Max_AUC": float(np.max(fold_auc)),
        "Mean_Of_Repeat_Means": float(np.mean(repeat_auc)),
        "SD_Of_Repeat_Means": float(np.std(repeat_auc, ddof=1)),
    }, repeat_means


def final_parameter_rows():
    rows = []
    for model, params in FINAL_PARAMS.items():
        for parameter, value in params.items():
            if isinstance(value, (list, tuple)):
                display_value = json.dumps(list(value))
            elif value is None:
                display_value = "None"
            else:
                display_value = value
            rows.append(
                {
                    "Model": model,
                    "Parameter": parameter,
                    "Final_Value": display_value,
                    "Fixed_For_Repeated_CV": True,
                }
            )
    return pd.DataFrame(rows)


def format_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True, name="Arial", size=10)
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet in workbook.worksheets:
        if sheet.max_row < 1:
            continue
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = border
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = border
        for column in sheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column]
            width = min(max(max((len(value) for value in values), default=8) + 2, 10), 42)
            sheet.column_dimensions[column[0].column_letter].width = width
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
    workbook.save(path)


def plot_repeated_results(fold_results, summary_results, output_dir):
    order = summary_results.sort_values(
        "Pooled_Fold_Mean_AUC", ascending=False
    )["Model"].tolist()

    plt.figure(figsize=(12, 7))
    sns.boxplot(
        data=fold_results,
        x="Model",
        y="AUC",
        order=order,
        color="#8FB9AA",
        width=0.6,
        fliersize=2,
    )
    sns.stripplot(
        data=fold_results,
        x="Model",
        y="AUC",
        order=order,
        color="#173F5F",
        size=2.5,
        alpha=0.6,
        jitter=0.18,
    )
    plt.axhline(0.5, color="#777777", linestyle="--", linewidth=1)
    plt.xlabel("Model")
    plt.ylabel("Fold-level AUC")
    plt.title("5x5 Repeated Stratified CV: Fold-level AUC")
    plt.xticks(rotation=30, ha="right")
    plt.tight_layout()
    plt.savefig(output_dir / "02_Fold_AUC_Distribution.svg", dpi=300)
    plt.close()

    repeat_means = (
        fold_results.groupby(["Model", "Repeat"], as_index=False)["AUC"].mean()
    )
    plt.figure(figsize=(12, 7))
    sns.lineplot(
        data=repeat_means,
        x="Repeat",
        y="AUC",
        hue="Model",
        marker="o",
        linewidth=1.8,
    )
    plt.axhline(0.5, color="#777777", linestyle="--", linewidth=1)
    plt.xticks([1, 2, 3, 4, 5])
    plt.xlabel("Repeated CV run")
    plt.ylabel("Mean AUC across five folds")
    plt.title("5x5 Repeated Stratified CV: Repeat-level AUC")
    plt.tight_layout()
    plt.savefig(output_dir / "03_Repeated_AUC_By_Run.svg", dpi=300)
    plt.close()


def run(args):
    base = load_base_pipeline()

    data, features, protected, _ = base.load_data(args.modality)
    full_train_df, internal_validation_df = train_test_split(
        data,
        test_size=0.2,
        random_state=base.RANDOM_STATE,
        stratify=data["label"],
    )
    full_train_df = full_train_df.reset_index(drop=True)
    internal_validation_df = internal_validation_df.reset_index(drop=True)

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    cv = RepeatedStratifiedKFold(
        n_splits=args.n_splits,
        n_repeats=args.n_repeats,
        random_state=args.random_state,
    )
    configs = base.model_configs()
    for model_name, params in FINAL_PARAMS.items():
        configs[model_name]["model"].set_params(**params)

    x_train_full = full_train_df[features].copy()
    y_train_full = full_train_df["label"].to_numpy()

    fold_rows = []
    fold_selection_rows = []
    total_folds = args.n_splits * args.n_repeats

    for split_index, (train_idx, test_idx) in enumerate(
        cv.split(x_train_full, y_train_full), start=1
    ):
        repeat = (split_index - 1) // args.n_splits + 1
        fold = (split_index - 1) % args.n_splits + 1
        print(
            f"Repeated CV {split_index:02d}/{total_folds}: "
            f"repeat {repeat}, fold {fold}"
        )

        x_fold_train = x_train_full.iloc[train_idx]
        x_fold_test = x_train_full.iloc[test_idx]
        y_fold_train = y_train_full[train_idx]
        y_fold_test = y_train_full[test_idx]

        selector = base.StrictPreprocessSelector(
            k=base.MRMR_K,
            protected_features=protected,
            missing_threshold=base.MISSING_THRESHOLD,
            variance_threshold=base.VARIANCE_THRESHOLD,
        )
        selector.fit(x_fold_train, y_fold_train)
        x_fold_train_selected = selector.transform(x_fold_train)
        x_fold_test_selected = selector.transform(x_fold_test)
        selected = list(selector.selected_features_)

        fold_selection_rows.append(
            {
                "Split_Index": split_index,
                "Repeat": repeat,
                "Fold": fold,
                "Train_N": len(train_idx),
                "Test_N": len(test_idx),
                "Selected_Count": len(selected),
                "Selected_Features": "; ".join(selected),
            }
        )

        for model_name, config in configs.items():
            classifier = clone(config["model"])
            classifier.fit(x_fold_train_selected, y_fold_train)
            probability = classifier.predict_proba(x_fold_test_selected)[:, 1]
            metrics = positive_rate_metrics(y_fold_test, probability)
            fold_rows.append(
                {
                    "Split_Index": split_index,
                    "Repeat": repeat,
                    "Fold": fold,
                    "Model": model_name,
                    "Train_N": len(train_idx),
                    "Test_N": len(test_idx),
                    "Selected_Count": len(selected),
                    **metrics,
                }
            )

    fold_results = pd.DataFrame(fold_rows)
    selection_results = pd.DataFrame(fold_selection_rows)

    summary_rows = []
    repeat_mean_frames = []
    for model_name in configs:
        model_fold = fold_results.loc[fold_results["Model"] == model_name].copy()
        summary, repeat_means = summarize_model_auc(model_fold)
        summary_rows.append(
            {
                "Model": model_name,
                "N_Folds": len(model_fold),
                **summary,
                "Mean_Fold_Brier": model_fold["Brier"].mean(),
                "Mean_Fold_Accuracy": model_fold["Accuracy"].mean(),
                "Mean_Fold_Recall": model_fold["Recall"].mean(),
                "Mean_Fold_Specificity": model_fold["Specificity"].mean(),
            }
        )
        repeat_means["Model"] = model_name
        repeat_mean_frames.append(repeat_means)

    summary_results = pd.DataFrame(summary_rows).sort_values(
        "Pooled_Fold_Mean_AUC", ascending=False
    )
    repeat_means_df = pd.concat(repeat_mean_frames, ignore_index=True)[
        ["Model", "Repeat", "Repeat_Mean_AUC"]
    ]

    parameter_df = final_parameter_rows()
    config_df = pd.DataFrame(
        [
            ["Input file", str(base.INPUT_FILE)],
            ["Analysis cohort", "Original training set only"],
            ["Training N", len(full_train_df)],
            ["Internal validation N (not used)", len(internal_validation_df)],
            ["Modality", args.modality],
            ["CV scheme", f"{args.n_splits}x{args.n_repeats} repeated stratified CV"],
            ["Number of fold evaluations per model", total_folds],
            ["Random seed", args.random_state],
            ["Classifier parameters", "Fixed to final strict-CV values"],
            ["Fold-internal preprocessing", "Missing filter -> mean imputation -> constant filter -> Z-score -> MWU/FDR -> mRMR"],
            ["Base pipeline", str(BASE_PIPELINE)],
            ["Run timestamp", datetime.now().isoformat(timespec="seconds")],
        ],
        columns=["Item", "Value"],
    )

    workbook_path = output_dir / "01_5x5_Repeated_CV_Results.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        summary_results.to_excel(writer, sheet_name="Summary", index=False)
        fold_results.to_excel(writer, sheet_name="Fold_AUCs", index=False)
        repeat_means_df.to_excel(writer, sheet_name="Repeat_Means", index=False)
        selection_results.to_excel(
            writer, sheet_name="Fold_Feature_Selection", index=False
        )
        parameter_df.to_excel(writer, sheet_name="Final_Parameters", index=False)
        config_df.to_excel(writer, sheet_name="CV_Config", index=False)
    format_workbook(workbook_path)

    summary_results.to_csv(output_dir / "04_Summary.csv", index=False, encoding="utf-8-sig")
    fold_results.to_csv(output_dir / "05_Fold_AUCs.csv", index=False, encoding="utf-8-sig")
    config = {
        "script": Path(__file__).name,
        "base_pipeline": str(BASE_PIPELINE),
        "input_file": str(base.INPUT_FILE),
        "modality": args.modality,
        "training_n": int(len(full_train_df)),
        "internal_validation_n_not_used": int(len(internal_validation_df)),
        "n_splits": args.n_splits,
        "n_repeats": args.n_repeats,
        "random_state": args.random_state,
        "final_params": FINAL_PARAMS,
        "output_workbook": str(workbook_path),
    }
    (output_dir / "06_Repeated_CV_Config.json").write_text(
        json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    plot_repeated_results(fold_results, summary_results, output_dir)

    print("\nSummary")
    print(
        summary_results[
            [
                "Model",
                "Pooled_Fold_Mean_AUC",
                "Pooled_Fold_SD",
                "Mean_Of_Repeat_Means",
                "SD_Of_Repeat_Means",
            ]
        ].to_string(index=False)
    )
    print(f"\nSaved results to: {output_dir}")


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--modality",
        choices=["combined", "clinical", "radiomics"],
        default="combined",
    )
    parser.add_argument("--n-splits", type=int, default=N_SPLITS)
    parser.add_argument("--n-repeats", type=int, default=N_REPEATS)
    parser.add_argument("--random-state", type=int, default=RANDOM_STATE)
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Defaults to 06e_StrictCV_折内预处理/09_5x5_Repeated_CV[_modality].",
    )
    args = parser.parse_args()
    if args.output_dir is None:
        suffix = "" if args.modality == "combined" else f"_{OUTPUT_LABELS[args.modality]}"
        args.output_dir = str(STRICT_ROOT / f"09_5x5_Repeated_CV{suffix}")
    return args


if __name__ == "__main__":
    run(parse_args())
